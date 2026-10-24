"""Line spacing as Word does it: the multiple times the font's own single line (Calibri's is
1.22 × its size), for English and Bangla text (Bangla letters count with the Bangla font's
line, so they don't run into the next line), smaller text in spans, a paragraph's own
multiple and an exact height, headings and tables, another Bangla font, lines on a page,
the Paragraph dialog and Line Spacing menu, Word export (multiples stay multiples, exact
heights stay exact) and import, undo, copy and save, the HTML export, and the font tables."""
import asyncio, sys, json, base64, zipfile, io, re
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)))
    if not cond: fails.append(what)

EN = 'The quick brown fox jumps over the lazy dog and keeps running across the field until evening comes. ' * 3
BN = 'আমি বাংলায় গান গাই, আমি বাংলার গান গাই। কৃষ্ণচূড়া ফুলে ফুলে রাঙা হয়ে উঠেছে পথের দুই ধার। ' * 3
HTML = (f'<p id="en">{EN}</p><p id="bn">{BN}</p><p id="small"><span style="font-size:9pt">{EN}</span></p>'
        f'<p id="m15" style="line-height: 1.5">{EN}</p><p id="exact" style="line-height: 14pt">{EN}</p>'
        f'<h1 id="h1">A heading long enough to run onto a second line of the page for sure, yes indeed</h1>'
        f'<table style="width:100%"><tbody><tr><td><p id="cen">{EN}</p></td><td><p id="cbn">{BN}</p></td></tr></tbody></table>'
        f'<p id="ns" class="nospacing">{EN}</p><p id="last">End.</p>')
# line pitch of a paragraph in CSS px, from its text's line boxes
PITCH = r"""(id) => { const p = document.getElementById(id); const r = document.createRange(); r.selectNodeContents(p);
  const tops = [...new Set([...r.getClientRects()].map(x => Math.round(x.top * 100) / 100))].sort((a, b) => a - b);
  const lines = []; for (const t of tops) if (!lines.length || t - lines[lines.length - 1] > 3) lines.push(t);
  return lines.length > 1 ? (lines[lines.length - 1] - lines[0]) / (lines.length - 1) / (__Q.S.zoom || 1) : null; }"""
PX = 96 / 72
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""

def word_docx():
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_LINE_SPACING
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.0; st.paragraph_format.space_after = Pt(0)
    d.add_paragraph(EN)
    p = d.add_paragraph(EN); p.paragraph_format.line_spacing = 1.5
    p = d.add_paragraph(EN); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(14)
    bio = io.BytesIO(); d.save(bio); return base64.b64encode(bio.getvalue()).decode()

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1300, height=1000, gfonts=True)
        await load(page)
        await page.evaluate("h => __Q.openDoc(__Q.newDocModel({ title: 'Lines', html: h }))", HTML)
        await page.evaluate("() => __Q.setZoom(1)")
        await page.wait_for_timeout(1500)
        st = lambda js, *a: page.evaluate(js, *a)
        pitch = lambda i: st(PITCH, i)
        fs = 11 * PX

        print('1. Word spacing')
        en = await pitch('en')
        ok(abs(en - 1.08 * 1.2207 * fs) < 0.1, 'Calibri 11 pt at the default 1.08: 1.08 × its single line of 1.2207 × 11 pt (14.5 pt, as Word)', en)
        bn = await pitch('bn')
        ok(abs(bn - 1.08 * 1.403 * fs) < 0.15, 'Bangla text counts with the Bangla font\'s line (Noto Sans Bengali, 1.403)', bn)
        ok(bn >= 1.325 * fs, 'so a Bangla line is taller than its letters (no overlap with the next line)', [bn, 1.325 * fs])
        sm = await pitch('small')
        ok(abs(sm - 1.08 * 1.2207 * 9 * PX) < 0.1, 'all text at 9 pt in an 11 pt paragraph gets 9 pt lines, as in Word', sm)
        m15 = await pitch('m15')
        ok(abs(m15 - 1.5 * 1.2207 * fs) < 0.1, 'a paragraph\'s own 1.5 is 1.5 × the single line', m15)
        ex = await pitch('exact')
        ok(abs(ex - 14 * PX) < 0.1, 'an exact 14 pt stays 14 pt', ex)
        h1 = await pitch('h1')
        ok(abs(h1 - 1.1 * 1.2207 * 16 * PX) < 0.15, 'Heading 1 (16 pt, 1.1)', h1)
        cen, cbn = await pitch('cen'), await pitch('cbn')
        ok(abs(cen - en) < 0.1 and abs(cbn - bn) < 0.15, 'the same in table cells', [cen, cbn])
        ns = await pitch('ns')
        ok(abs(ns - 1.2207 * fs) < 0.1, 'No Spacing is single (1.0)', ns)
        marks = await st("() => ['en', 'bn', 'small', 'm15', 'exact'].map(i => document.getElementById(i).getAttribute('data-lh'))")
        ok(marks[0] is None and marks[1] and marks[2] and marks[2].endswith('z') and marks[3] and marks[4] is None, 'only paragraphs that differ from their style are marked', marks)

        print('2. the dialog, the menu and typing')
        info = await st("id => { const p = document.getElementById(id); const r = document.createRange(); r.setStart(p.firstChild, 3); r.collapse(true); getSelection().removeAllRanges(); getSelection().addRange(r); return __Q.Dialogs && window.__Q.UI ? null : null; }", 'm15')
        line = await st("() => { const cs = getComputedStyle(document.getElementById('m15')); return parseFloat(cs.getPropertyValue('--lhm')); }")
        para = await st("() => { const b = document.getElementById('m15'); const r = document.createRange(); r.setStart(b.firstChild, 3); r.collapse(true); getSelection().removeAllRanges(); getSelection().addRange(r); return [document.getElementById('m15').style.lineHeight]; }")
        await page.wait_for_timeout(100)
        await st("() => __Q.UI.selectTab('home')")
        await page.locator('#ribbon .rpanel.active [data-cmd="lineSpacing"], #ribbon .rpanel.active button[aria-label^="Line and Paragraph Spacing"]').first.click()
        await page.wait_for_timeout(250)
        checked = await st("() => [...document.querySelectorAll('.pop .mi')].filter(m => m.getAttribute('aria-checked') === 'true' || m.classList.contains('on') || m.querySelector('.ck, .check')).map(m => m.textContent.trim())")
        ok(any(c.startswith('1.5') for c in checked), 'the Line Spacing menu shows 1.5 for that paragraph', checked)
        await page.keyboard.press('Escape')
        await st("() => { const b = document.getElementById('en'); const r = document.createRange(); r.setStart(b.firstChild, 3); r.collapse(true); getSelection().removeAllRanges(); getSelection().addRange(r); }")
        await page.wait_for_timeout(100)
        await page.locator('#ribbon .rpanel.active [data-cmd="lineSpacing"], #ribbon .rpanel.active button[aria-label^="Line and Paragraph Spacing"]').first.click()
        await page.wait_for_timeout(250)
        checked2 = await st("() => [...document.querySelectorAll('.pop .mi')].filter(m => m.getAttribute('aria-checked') === 'true' || m.classList.contains('on') || m.querySelector('.ck, .check')).map(m => m.textContent.trim())")
        ok(any(c.startswith('1.08') for c in checked2), 'and 1.08 for a Normal paragraph (the Word value, not the CSS factor)', checked2)
        await page.keyboard.press('Escape')
        # typing Bangla into an English paragraph gives it room at once
        await st("() => { const b = document.getElementById('last'); const r = document.createRange(); r.setStart(b.firstChild, 4); r.collapse(true); getSelection().removeAllRanges(); getSelection().addRange(r); }")
        await page.keyboard.insert_text(' এবং আরও কিছু বাংলা লেখা যা এক লাইনের বেশি হবে, কারণ এখানে অনেক শব্দ আছে এবং আরও আছে অনেক অনেক শব্দ। ' * 2)
        await page.wait_for_timeout(200)
        lp = await pitch('last')
        ok(abs(lp - bn) < 0.2, 'typing Bangla into a paragraph spaces it for Bangla right away', [lp, bn])
        await st("() => __Q.undo()")
        mk = await st("() => document.getElementById('bn').getAttribute('data-lh')")
        ok(mk == marks[1], 'after undo the marks are back at once', mk)

        print('3. a page of lines')
        await st("t => { __Q.openDoc(__Q.newDocModel({ title: 'Page', html: '<p id=\"one\" style=\"margin-bottom:0\">' + t + '</p>' })); __Q.setZoom(1); }", EN * 40)
        await page.wait_for_timeout(800)
        n1 = await st("""() => { const g = __Q.geom(), pitch = g.h + __Q.GAP; const p = document.getElementById('one'); const r = document.createRange(); r.selectNodeContents(p);
          const top0 = document.getElementById('editor').getBoundingClientRect().top; const tops = new Set();
          for (const q of r.getClientRects()) { const y = q.top - top0; if (y < pitch && q.height > 5 && q.height < 40 && q.width > 1) tops.add(Math.round(y)); } return [tops.size, __Q.geom().h, __Q.geom().mt, __Q.geom().mb]; }""")
        ok(n1[0] == 48, 'an A4 page with 1" margins holds 48 lines of Calibri 11 at 1.08, as in Word', n1)

        print('4. Word files')
        await page.evaluate("h => { __Q.openDoc(__Q.newDocModel({ title: 'Lines', html: h })); __Q.setZoom(1); }", HTML)
        await page.wait_for_timeout(800)
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        styles = z.read('word/styles.xml').decode(); body = z.read('word/document.xml').decode()
        def style_line(sid):
            m = re.search(r'w:styleId="%s".*?</w:style>' % sid, styles, re.S)
            l = re.search(r'w:line="(\d+)"', m.group(0)) if m else None
            return int(l.group(1)) if l else None
        dl = re.search(r'<w:pPrDefault>.*?w:line="(\d+)"', styles, re.S)
        ok(dl and int(dl.group(1)) == 259, 'the default paragraph spacing is Word\'s 1.08 (259)', dl and dl.group(1))
        got = {k: style_line(k) for k in ['NoSpacing', 'Heading1', 'Title', 'Code']}
        ok(got == {'NoSpacing': 240, 'Heading1': 264, 'Title': 252, 'Code': 264}, 'styles keep their multiples (No Spacing 1, headings 1.1, Title 1.05, Code 1.1)', got)
        ok('w:line="360" w:lineRule="auto"' in body, 'a paragraph at 1.5 is 360 (auto)')
        ok('w:line="280" w:lineRule="exact"' in body, 'an exact 14 pt stays exact (280 twips)')
        ok('data-lh' not in await st("() => __Q.cleanHTML()"), 'marks are not saved')
        await st("() => { const r = document.createRange(); r.selectNodeContents(document.getElementById('bn')); getSelection().removeAllRanges(); getSelection().addRange(r); }")
        html = await st("() => { const dt = new DataTransfer(); const e = new ClipboardEvent('copy', { clipboardData: dt, bubbles: true, cancelable: true }); document.getElementById('editor').dispatchEvent(e); return dt.getData('text/html'); }")
        ok('data-lh' not in html, 'nor copied')
        sa = await st("() => __Q.toStandaloneHtml()")
        ok('--doc-r:1.2207' in sa and 'data-lh="1.5152"' in sa and '[data-lh="1.5152"]' in sa, 'the HTML export carries the spacing (font line and marks with their rules)')
        # import: Normal at single, a paragraph at 1.5, one exactly 14 pt
        r = await st("""async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i);
          const r = await __Q.readDocxFile(u.buffer); await __Q.openImported(r); __Q.setZoom(1); return [__Q.S.doc.line, r.html.slice(0, 200)]; }""", word_docx())
        await page.wait_for_timeout(900)
        ok(r[0] == 1, 'a Word file with Normal at single opens with single spacing', r)
        ps = await st("() => [...document.querySelectorAll('#editor > p')].slice(0, 3).map((p, i) => { p.id = 'w' + i; return p.style.lineHeight; })")
        pw = [await pitch('w0'), await pitch('w1'), await pitch('w2')]
        ok(abs(pw[0] - 1.2207 * fs) < 0.1 and abs(pw[1] - 1.5 * 1.2207 * fs) < 0.1 and abs(pw[2] - 14 * PX) < 0.1, 'its paragraphs space like Word: single, 1.5 and exactly 14 pt', [ps, pw])

        print('5. fonts')
        rat = await st("""async () => { const out = {}; for (const f of ['carlito-r', 'notosansbengali-r', 'arimo-r', 'tinos-r']) { const b = await (await fetch('fonts/' + f + '.ttf')).arrayBuffer(); out[f] = Math.round(__Q.wordLineOf(b) * 10000) / 10000; } return out; }""")
        ok(rat == {'carlito-r': 1.2207, 'notosansbengali-r': 1.403, 'arimo-r': 1.4321, 'tinos-r': 1.1499}, 'Word\'s single line read from font files (winAscent + winDescent + gap)', rat)
        ok(await st("() => __Q.LineH.family('Arial')") == 1.1499, 'Arial is spaced as Arial in Word, not as its look-alike')
        await page.evaluate("h => { __Q.openDoc(__Q.newDocModel({ title: 'Lines', html: h })); __Q.setZoom(1); }", HTML)
        await page.wait_for_timeout(600)
        await st("() => __Q.setDocBnFont('Hind Siliguri')")
        await page.wait_for_timeout(900)
        hs = await pitch('bn')
        ok(abs(hs - 1.08 * 1.617 * fs) < 0.2, 'with Hind Siliguri as the Bangla font, Bangla lines follow its line (1.617)', hs)
        ok(abs(await pitch('en') - en) < 0.05, 'English lines stay as they were', await pitch('en'))
        ok(not [e for e in errors if e[0] == 'pageerror'], 'no page errors', errors)
        await b.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')

asyncio.run(main())
