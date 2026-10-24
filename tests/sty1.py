"""Styles and multilevel lists. Modify a built-in style (Heading 1, Normal) in the Modify Style
dialog and every paragraph with it follows (undo too); Create a Style from a formatted paragraph
(its own formatting moves into the style), apply it from the gallery, the Styles pane and the
All Styles menu; Update to Match Selection; the style for the following paragraph on Enter; rename,
select all, delete, reset. Lists: the Bangla and bracketed numbers of the Numbering Library; the
Multilevel List presets (1.1.1, ১। (ক) (১), ১.১, (1) (a) (i), I. A. 1.); on screen, in the PDF,
in Word files out and back (styles.xml, numbering.xml with Word 2010's custom Bangla formats), a
python-docx file with a style of its own, LibreOffice; the web page; phone width."""
import asyncio, sys, json, base64, io, zipfile, re, subprocess, os, tempfile
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)[:700]))
    if not cond: fails.append(what)

HTML = ('<h1>Office order</h1><p>The office will remain closed on Sunday for the national holiday.</p>'
        '<p style="text-align:justify;margin-left:36pt"><span style="font-family:Georgia;font-size:12pt;color:#1f3a5f">All staff are requested to note this and plan their work accordingly.</span></p>'
        '<h1>Arrangements</h1><p>Duty officers will be named in a separate order.</p>'
        '<ol><li>Accounts<ol><li>Salary bills</li><li>Pending vouchers<ol><li>Over one month old</li></ol></li></ol></li><li>Administration</li></ol>'
        '<p>Copy to all sections.</p><p>End of order.</p>')
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""
OPEN_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); r.title = 'Imported'; await __Q.openImported(r); return true; }"""
CARET = """([sel, i, off]) => { const el = document.querySelectorAll('#editor ' + sel)[i]; const w = document.createTreeWalker(el, NodeFilter.SHOW_TEXT); const t = w.nextNode(); const r = document.createRange(); r.setStart(t, Math.min(off, t.length)); r.collapse(true); getSelection().removeAllRanges(); getSelection().addRange(r); document.getElementById('editor').focus(); }"""
LOOK = """sel => [...document.querySelectorAll('#editor ' + sel)].map(e => { const cs = getComputedStyle(e); return [cs.fontFamily.split(',')[0].replace(/"/g, ''), cs.fontSize, cs.color, cs.textAlign, cs.marginLeft]; })"""
MARKS = """() => [...document.querySelectorAll('#editor li')].map(li => __Q.markerString(li, getComputedStyle(li).listStyleType).trim())"""

def py_docx():
    """a Word file with a paragraph style of its own and a Bangla list"""
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    st = d.styles.add_style('Office Body', WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = d.styles['Normal']
    st.font.name = 'Georgia'; st.font.size = Pt(12); st.font.color.rgb = RGBColor(0x1f, 0x3a, 0x5f)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    st.paragraph_format.space_after = Pt(6)
    h1 = d.styles['Heading 1']
    h1.font.name = 'Times New Roman'; h1.font.size = Pt(18); h1.font.color.rgb = RGBColor(0x80, 0, 0); h1.font.bold = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_heading('Office order', level=1)
    d.add_paragraph('A paragraph in the Office Body style.', style='Office Body')
    d.add_paragraph('Another one, with a bold word.', style='Office Body').runs[0].bold = False
    d.add_paragraph('A plain Normal paragraph, justified as Normal says.')
    d.add_paragraph('One aligned left on its own.').alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_heading('Arrangements', level=1)
    bio = io.BytesIO(); d.save(bio); return base64.b64encode(bio.getvalue()).decode()

def pdf_text(path):
    return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1366, height=900, mock_caps=True, gfonts=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(500)

        print('1. Modify Heading 1 (the dialog)')
        await st(CARET, ['h1', 0, 2])
        await st("() => __Q.UI.selectTab('home')")
        g = await st("() => [...document.querySelectorAll('#styleGalleryItems .gal-item')].map(x => x.dataset.style)")
        ok(g[:7] == ['normal', 'nospacing', 'title', 'subtitle', 'h1', 'h2', 'h3'], 'the gallery shows Word\'s styles', g)
        await page.locator('#styleGalleryItems .gal-item[data-style="h1"]').click(button='right')
        await page.wait_for_timeout(200)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(m => m.textContent.trim())")
        ok(any(i.startswith('Update Heading 1 to Match Selection') for i in items) and any(i.startswith('Modify') for i in items) and any('Select All 2 Instances' in i for i in items), 'right-click on a style: Update to Match, Modify…, Select All 2 Instances', items)
        await page.locator('.pop .mi', has_text='Modify').click()
        await page.wait_for_timeout(300)
        dlg = page.locator('.dialog')
        await dlg.locator('select').nth(1).select_option('Times New Roman')
        await dlg.locator('input[type=number]').nth(0).fill('18')
        await dlg.locator('input[type=color]').evaluate("(c) => { c.value = '#000000'; c.dispatchEvent(new Event('input')); }")
        await dlg.locator('select').nth(2).select_option('center')
        prev = await st("() => getComputedStyle(document.querySelector('.st-sample p')).textAlign")
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        look = await st(LOOK, 'h1')
        ok(prev == 'center' and all(l[0] == 'Times New Roman' and l[1] == '24px' and l[2] == 'rgb(0, 0, 0)' and l[3] == 'center' for l in look), 'both headings change: Times New Roman 18 pt, black, centred (the preview showed it)', look)
        ok(await st("() => __Q.S.doc.styles.h1") == {'font': 'Times New Roman', 'size': 18, 'color': '#000000', 'align': 'center'}, 'only what changed is kept', await st("() => __Q.S.doc.styles.h1"))
        await st("() => __Q.undo()")
        await page.wait_for_timeout(250)
        ok((await st(LOOK, 'h1'))[0][0] == 'Calibri' and not await st("() => __Q.S.doc.styles.h1"), 'Undo puts the style back', await st(LOOK, 'h1'))
        await st("() => __Q.redo()")
        await page.wait_for_timeout(250)
        ok((await st(LOOK, 'h1'))[1][0] == 'Times New Roman', 'Redo changes it again')

        print('2. Create a Style from a paragraph; apply it')
        await st(CARET, ['p span', 0, 5])
        await page.locator('#styleGalleryItems').locator('..').locator('.gallery-nav button').nth(2).click()
        await page.wait_for_timeout(200)
        await page.locator('.pop .mi', has_text='Create a Style').click()
        await page.wait_for_timeout(300)
        name_in = page.locator('.dialog input').first
        await name_in.fill('Office Body')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        r = await st("() => { const k = __Q.Styles.custom()[0]; return { k, def: __Q.S.doc.styles[k], html: document.querySelectorAll('#editor > p')[1].outerHTML }; }")
        k = r['k']
        ok(r['def'] and r['def']['name'] == 'Office Body' and r['def'].get('font') == 'Georgia' and r['def'].get('size') == 12 and r['def'].get('color') == '#1f3a5f' and r['def'].get('align') == 'justify' and r['def'].get('indent') == 36, 'the style takes the paragraph\'s font, size, colour, alignment and indent', r['def'])
        ok(r['html'] == f'<p class="{k}">All staff are requested to note this and plan their work accordingly.</p>', 'the paragraph has the style, its own formatting moved into it', r['html'])
        await st(CARET, ['p', 0, 3])
        await page.locator(f'#styleGalleryItems .gal-item[data-style="{k}"]').click()
        await page.wait_for_timeout(250)
        look = await st(LOOK, f'p.{k}')
        ok(len(look) == 2 and all(l[0] == 'Georgia' and l[1] == '16px' and l[3] == 'justify' and l[4] == '48px' for l in look), 'clicking it in the gallery gives another paragraph the same look', look)

        print('3. Update to Match Selection; Enter; Normal')
        await st("k => __Q.change(() => { const p = document.querySelector('#editor p.' + k); p.style.fontSize = '14pt'; p.style.textAlign = 'left'; })", k)
        await st(CARET, [f'p.{k}', 0, 2])
        await st("k => __Q.Styles.update(k)", k)
        await page.wait_for_timeout(300)
        look = await st(LOOK, f'p.{k}')
        ok(all(l[1] == '18.6667px' and l[3] in ('left', 'start') for l in look) and await st("k => document.querySelector('#editor p.' + k).getAttribute('style')", k) is None, 'Update to Match: both paragraphs take 14 pt, left; the first one\'s own formatting goes into the style', look)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(250)
        r = await st("k => [document.querySelector('#editor p.' + k).getAttribute('style'), __Q.S.doc.styles[k].size, __Q.S.doc.styles[k].align]", k)
        ok(r == ['font-size: 14pt; text-align: left;', 12, 'justify'], 'one Undo: the style and the paragraph as they were', r)
        await st("() => __Q.redo()")
        await page.wait_for_timeout(250)
        # Enter at the end: the style goes on (its own next); set next to Normal and try again
        await st("k => { const p = document.querySelectorAll('#editor p.' + k)[1]; const r = document.createRange(); r.selectNodeContents(p); r.collapse(false); getSelection().removeAllRanges(); getSelection().addRange(r); }", k)
        await page.keyboard.press('Enter')
        await page.wait_for_timeout(150)
        cls1 = await st("() => { const r = getSelection().getRangeAt(0); const b = r.startContainer.nodeType === 1 ? r.startContainer : r.startContainer.parentNode; return b.closest('p').className; }")
        await page.keyboard.press('Backspace')
        await st("k => { __Q.S.doc.styles[k].next = 'normal'; }", k)
        await st("k => { const p = document.querySelectorAll('#editor p.' + k)[1]; const r = document.createRange(); r.selectNodeContents(p); r.collapse(false); getSelection().removeAllRanges(); getSelection().addRange(r); }", k)
        await page.keyboard.press('Enter')
        await page.wait_for_timeout(150)
        cls2 = await st("() => { const r = getSelection().getRangeAt(0); const b = r.startContainer.nodeType === 1 ? r.startContainer : r.startContainer.parentNode; return b.closest('p').className; }")
        await page.keyboard.press('Backspace')
        ok(cls1 == k and cls2 == '', 'Enter at its end: the same style follows; with Normal as the next style, a Normal paragraph', [cls1, cls2])
        # Normal: justified, Times New Roman 12; the document's style starts from Normal
        await st("() => __Q.Styles.modify('normal', Object.assign(__Q.Styles.look('normal'), { font: 'Times New Roman', size: 12, align: 'justify' }))")
        await page.wait_for_timeout(300)
        r = await st("() => ({ font: __Q.S.doc.font, size: __Q.S.doc.size, normal: __Q.S.doc.styles.normal, look: [...document.querySelectorAll('#editor > p:not([class])')].map(p => getComputedStyle(p).textAlign + ':' + getComputedStyle(p).fontFamily.split(',')[0]) })")
        ok(r['font'] == 'Times New Roman' and r['size'] == 12 and r['normal'] == {'align': 'justify'} and all(x.startswith('justify:"Times New Roman"') for x in r['look']), 'Modify Normal: the document\'s font and size, plain paragraphs justified', r)

        print('4. The Styles pane; rename, select all, delete, reset')
        await page.locator('.rgroup[aria-label="Styles"] .launcher').click()
        await page.wait_for_timeout(300)
        rows = await st("() => [...document.querySelectorAll('.st-item')].map(x => [x.dataset.style, x.querySelector('.st-nm').textContent, x.querySelector('.st-n').textContent, x.classList.contains('on')])")
        ok(any(x[0] == k and x[1] == 'Office Body' and x[2] == '2' for x in rows) and any(x[0] == 'h1' and x[2] == '2' for x in rows), 'the Styles pane lists the styles with how many paragraphs have each', rows)
        await st(CARET, ['p', 0, 1])
        await page.locator('.st-item[data-style="h2"]').click()
        await page.wait_for_timeout(250)
        ok(await st("() => document.querySelector('#editor h2') && document.querySelector('#editor h2').textContent.startsWith('The office')") and await st("() => document.querySelector('.st-item.on').dataset.style") == 'h2', 'a click in the pane applies the style (and marks it as the current one)')
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        await page.locator(f'.st-item[data-style="{k}"] .st-more').click()
        await page.wait_for_timeout(200)
        await page.locator('.pop .mi', has_text='Rename').click()
        await page.wait_for_timeout(200)
        await page.locator('.dialog input').first.fill('Office Text')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(250)
        ok(await st("k => __Q.Styles.name(k)", k) == 'Office Text' and await st("k => document.querySelector(`#styleGalleryItems .gal-item[data-style='${k}'] .nm`).textContent", k) == 'Office Text', 'Rename: the gallery shows the new name')
        await st("k => __Q.Styles.selectAll(k)", k)
        r = await st("k => { const s = getSelection().getRangeAt(0); const ps = document.querySelectorAll('#editor p.' + k); return [ps[0].contains(s.startContainer) || s.startContainer === ps[0], ps[1].contains(s.endContainer) || s.endContainer === ps[1]]; }", k)
        ok(r == [True, True], 'Select All Instances: from the first paragraph with it to the last', r)
        await st("() => __Q.Styles.reset('h1')")
        await page.wait_for_timeout(250)
        r = await st(LOOK, 'h1')
        ok(r[0][1] == '21.3333px' and r[0][2] == 'rgb(31, 58, 95)' and r[0][3] == 'start' and not await st("() => __Q.S.doc.styles.h1"), 'Reset: Heading 1 looks as Folio draws it again (16 pt, blue, left)', r)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(250)

        print('5. Lists: Numbering Library, Multilevel List')
        await st(CARET, ['li', 0, 1])
        await st("() => __Q.CMD.numbering('bn-dari')")
        await page.wait_for_timeout(200)
        m = await st(MARKS)
        ok(m[0] == '১।' and m[-1] == '২।', 'Numbering ১। ২। ৩।', m)
        await page.locator('.rb:visible[title^="Multilevel List"]').click()
        await page.wait_for_timeout(200)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(m => m.textContent.trim())")
        ok(any('1.1.1.' in i for i in items) and any('(ক)' in i for i in items) and any('১.১.' in i for i in items), 'the Multilevel List menu has 1.1.1., ১। (ক) (১), ১.১.১. and more', items)
        await page.locator('.pop .mi', has_text='1.  1.1.  1.1.1.').click()
        await page.wait_for_timeout(250)
        m = await st(MARKS)
        ok(m == ['1.', '1.1.', '1.2.', '1.2.1.', '2.'], 'legal numbering: 1. 1.1. 1.2. 1.2.1. 2.', m)
        ok(await st("() => document.querySelector('#editor ol').getAttribute('data-ml')") == 'legal' and await st("() => document.querySelector('#editor ol').getAttribute('style')") is None, 'the list\'s own number type gave way to the preset')
        await st("() => __Q.Lists.multilevel('bn')")
        await page.wait_for_timeout(200)
        ok(await st(MARKS) == ['১।', '(ক)', '(খ)', '(১)', '২।'], 'Bangla office style: ১। (ক) (১)', await st(MARKS))
        await st("() => __Q.Lists.multilevel('bn-legal')")
        ok(await st(MARKS) == ['১.', '১.১.', '১.২.', '১.২.১.', '২.'], '১. ১.১. ১.২.১.', await st(MARKS))
        await st("() => __Q.Lists.multilevel('paren')")
        ok(await st(MARKS) == ['(1)', '(a)', '(b)', '(i)', '(2)'], '(1) (a) (i)', await st(MARKS))
        await st("() => __Q.Lists.multilevel('outline')")
        ok(await st(MARKS) == ['I.', 'A.', 'B.', '1.', 'II.'], 'I. A. 1.', await st(MARKS))
        await st("() => __Q.Lists.multilevel('legal')")
        # a plain paragraph made a list with a preset
        n0 = await st("() => document.querySelectorAll('#editor ol').length")
        await st(CARET, ['p', 4, 2])
        await st("() => __Q.Lists.multilevel('bn')")
        await page.wait_for_timeout(200)
        r = await st("() => { const r = getSelection().getRangeAt(0); const li = (r.startContainer.nodeType === 1 ? r.startContainer : r.startContainer.parentNode).closest('li'); return li && [li.parentElement.getAttribute('data-ml'), __Q.markerString(li, getComputedStyle(li).listStyleType), li.textContent]; }")
        ok(r and r[0] == 'bn' and r[1].strip() == '১।' and r[2] == 'End of order.', 'a paragraph becomes a list numbered with the preset', r)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        r = await st("() => [document.querySelectorAll('#editor ol').length, [...document.querySelectorAll('#editor > p')].pop().outerHTML]")
        ok(r == [n0, '<p>End of order.</p>'], 'one Undo: the paragraph again', r)

        print('6. The Word file')
        await st("() => __Q.CMD.numbering && 0")
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        sty = z.read('word/styles.xml').decode()
        doc = z.read('word/document.xml').decode()
        num = z.read('word/numbering.xml').decode()
        h1 = re.search(r'<w:style w:type="paragraph" w:styleId="Heading1">.*?</w:style>', sty).group(0)
        dd = re.search(r'<w:docDefaults>.*?</w:docDefaults>', sty).group(0)
        ok(('w:ascii="Times New Roman"' in h1 or ('<w:rFonts' not in h1 and 'w:ascii="Times New Roman"' in dd)) and '<w:sz w:val="36"/>' in h1 and '<w:jc w:val="center"/>' in h1 and '<w:color' not in h1.split('<w:rPr>')[1], 'Heading 1 in styles.xml: Times New Roman 18 pt (the document\'s font), centred, black', h1)
        own = re.search(r'<w:style w:type="paragraph" w:customStyle="1" w:styleId="OfficeText">.*?</w:style>', sty)
        ok(own and '<w:name w:val="Office Text"/>' in own.group(0) and '<w:basedOn w:val="Normal"/>' in own.group(0) and '<w:next w:val="Normal"/>' in own.group(0) and 'Georgia' in own.group(0) and '<w:sz w:val="28"/>' in own.group(0) and 'w:left="720"' in own.group(0), 'the document\'s style: Office Text, based on Normal, next Normal, Georgia 14 pt, indented', own and own.group(0))
        paras = re.findall(r'<w:p>(?:(?!</w:p>).)*?<w:pStyle w:val="OfficeText"/>(?:(?!</w:p>).)*?</w:p>', doc)
        ok(len(paras) == 2 and not any('<w:jc ' in x or '<w:ind ' in x or '<w:spacing ' in x or 'Georgia' in x for x in paras), 'its paragraphs name the style and repeat none of it', paras[:1])
        nd = re.search(r'<w:docDefaults>.*?</w:docDefaults>', sty).group(0)
        ok('w:ascii="Times New Roman"' in nd and '<w:sz w:val="24"/>' in nd and '<w:jc w:val="both"/>' in re.search(r'w:styleId="Normal">.*?</w:style>', sty).group(0), 'Normal: Times New Roman 12 pt in the defaults, justified', nd)
        ok('xmlns:mc=' in num and 'mc:Ignorable="w14 wp14"' in num and '<w:lvlText w:val="%1.%2."/>' in num and '<w:lvlText w:val="%1.%2.%3."/>' in num, 'numbering.xml: the legal list numbers each level %1.%2.', num[:300])
        # Bangla formats
        await st(CARET, ['li', 0, 1])
        await st("() => __Q.Lists.multilevel('bn')")
        b64b = await st(GET_DOCX)
        numb = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64b))).read('word/numbering.xml').decode()
        ok('<mc:Choice Requires="w14 wp14"><w:numFmt w:val="custom" w:format="১, ২, ৩, ..."/></mc:Choice><mc:Fallback><w:numFmt w:val="decimal"/></mc:Fallback>' in numb and 'xmlns:wp14=' in numb and '<w:lvlText w:val="%1।"/>' in numb and 'w:format="ক, খ, গ, ..."' in numb and '<w:lvlText w:val="(%2)"/>' in numb, 'Bangla numbers as Word 2010 custom formats (decimal for older readers): %1। and (%2) with ক, খ, গ', re.findall(r'<w:lvl w:ilvl="[01]">.*?</w:lvl>', numb)[:2])
        await st(OPEN_DOCX, b64b)
        await page.wait_for_timeout(700)
        r = await st("() => ({ styles: __Q.S.doc.styles, marks: [...document.querySelectorAll('#editor li')].map(li => __Q.markerString(li, getComputedStyle(li).listStyleType).trim()), own: [...document.querySelectorAll('#editor p[class^=\"s-\"]')].map(p => p.outerHTML) })")
        ownk = [x for x in r['styles'] if x.startswith('s-')]
        ok(len(ownk) == 1 and r['styles'][ownk[0]]['name'] == 'Office Text' and r['styles'][ownk[0]].get('font') == 'Georgia' and r['styles'][ownk[0]].get('size') == 14 and r['styles'][ownk[0]].get('next') == 'normal', 'opened again: Office Text is a document style (font, size, next)', r['styles'])
        ok(len(r['own']) == 2 and all('style=' not in x and '<span' not in x for x in r['own']), 'its paragraphs have the style and nothing of their own', r['own'])
        ok(r['marks'] == ['১।', '(ক)', '(খ)', '(১)', '২।'], 'the Bangla list numbers come back', r['marks'])
        await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(700)
        ok(await st(MARKS) == ['1.', '1.1.', '1.2.', '1.2.1.', '2.'] and await st("() => document.querySelector('#editor ol').getAttribute('data-ml')") == 'legal', 'the legal list comes back as a Multilevel List', await st(MARKS))
        # a Word file with a style of its own
        await st(OPEN_DOCX, py_docx())
        await page.wait_for_timeout(700)
        r = await st("() => ({ styles: __Q.S.doc.styles, ps: [...document.querySelectorAll('#editor p')].map(p => [p.className, p.getAttribute('style'), getComputedStyle(p).fontFamily.split(',')[0], getComputedStyle(p).textAlign]) })")
        ownk = [x for x in (r['styles'] or {}) if x.startswith('s-')]
        ok(len(ownk) == 1 and r['styles'][ownk[0]]['name'] == 'Office Body' and r['styles'][ownk[0]].get('align') == 'justify' and r['styles'][ownk[0]].get('after') == 6, 'a python-docx style becomes a document style', r['styles'])
        ok(r['ps'][0][0] == ownk[0] and r['ps'][0][2] == 'Georgia' and r['ps'][0][3] == 'justify' and r['ps'][2][0] == '', 'its paragraphs look as in Word, the Normal one plain', r['ps'])
        ok(r['styles'].get('normal') == {'align': 'justify'} and r['ps'][2][1] is None and r['ps'][2][3] == 'justify' and r['ps'][3][1] == 'text-align:left' and r['ps'][3][3] == 'left', 'Word\'s justified Normal is the Normal style: plain paragraphs say nothing, the left one does', [r['styles'].get('normal'), r['ps'][2:]])
        h = r['styles'].get('h1') or {}
        hl = await st("() => [...document.querySelectorAll('#editor h1')].map(e => [e.getAttribute('style'), getComputedStyle(e).fontFamily.split(',')[0].replace(/\"/g, ''), getComputedStyle(e).fontSize, getComputedStyle(e).color, getComputedStyle(e).textAlign])")
        ok(h.get('font') == 'Times New Roman' and h.get('size') == 18 and h.get('color') == '#800000' and h.get('align') == 'center' and 'bold' not in h, 'Word\'s Heading 1 is the document\'s Heading 1 (what differs from Folio\'s)', h)
        ok(len(hl) == 2 and all(x[0] is None and x[1] == 'Times New Roman' and x[2] == '24px' and x[3] == 'rgb(128, 0, 0)' and x[4] == 'center' for x in hl), 'the headings carry nothing of their own', hl)
        await st("() => __Q.Styles.modify('h1', Object.assign(__Q.Styles.look('h1'), { color: '#000000' }))")
        await page.wait_for_timeout(250)
        ok(await st("() => [...document.querySelectorAll('#editor h1')].every(e => getComputedStyle(e).color === 'rgb(0, 0, 0)')"), 'so Modify Heading 1 changes them all')
        ok(await st("() => __Q.Hist.idx") == 1, 'opening the file is no step to undo', await st("() => __Q.Hist.idx"))
        # LibreOffice reads ours
        with tempfile.TemporaryDirectory() as td:
            for nm, data in (('legal', b64), ('bn', b64b)):
                src = os.path.join(td, nm + '.docx'); open(src, 'wb').write(base64.b64decode(data))
                subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', td, src], capture_output=True, timeout=120)
            t1 = pdf_text(os.path.join(td, 'legal.pdf')) if os.path.exists(os.path.join(td, 'legal.pdf')) else ''
            t2 = pdf_text(os.path.join(td, 'bn.pdf')) if os.path.exists(os.path.join(td, 'bn.pdf')) else ''
            ok('1.2.1.' in t1 and 'Office order' in t1, 'LibreOffice: the legal numbers', t1[:500])
            ok(re.search(r'(১।|1।) Accounts', t2) is not None and re.search(r'\((ক|a)\) Salary', t2) is not None and re.search(r'\((১|1)\) Over one', t2) is not None, 'LibreOffice: the Bangla list keeps its shape (১। or 1।, (ক) or (a))', t2[:600])

        print('7. PDF, web page, phone')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(500)
        r = await st("() => [__Q.S.pane, [...document.querySelectorAll('.st-item')].map(x => x.dataset.style + ':' + x.querySelector('.st-n').textContent)]")
        ok(r[0] == 'styles' and not any(x.startswith('s-') for x in r[1]) and 'normal:5' in r[1] and 'h1:2' in r[1], 'the Styles pane shows the document just opened', r)
        await st("() => __Q.Styles.modify('h1', Object.assign(__Q.Styles.look('h1'), { font: 'Times New Roman', size: 18 }))")
        await st(CARET, ['li', 0, 1])
        await st("() => __Q.Lists.multilevel('legal')")
        await st(CARET, ['li', 4, 1])
        await st("() => __Q.CMD.numbering('bn-alpha-paren')")
        await page.wait_for_timeout(400)
        r = await st("""async () => { await document.fonts.ready; const blob = await __Q.buildVectorPdf({ title: 'Styles' }); const u = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < u.length; i += 0x8000) s += String.fromCharCode.apply(null, u.subarray(i, i + 0x8000)); return btoa(s); }""")
        open('shots/sty1.pdf', 'wb').write(base64.b64decode(r))
        txt = pdf_text('shots/sty1.pdf')
        fonts = subprocess.run(['pdffonts', 'shots/sty1.pdf'], capture_output=True, text=True).stdout
        ok('1.2.1.' in txt and '1.1.' in txt, 'PDF: the legal numbers are drawn', txt[:400])
        ok('Tinos' in fonts or 'Times' in fonts, 'PDF: the heading in Times New Roman (its free look-alike)', fonts)
        html = await st("() => __Q.toStandaloneHtml()")
        ok('.doc h1{font-family:' in html.replace(' ', '') or re.search(r'\.doc h1\{font-family', html) is not None, 'the web page carries the changed style', re.findall(r'\.doc h1\{[^}]*\}', html)[:1])
        b2, page2, errors2 = await new_page(p, width=400, height=860, mock_caps=True, gfonts=True)
        await load(page2)
        await page2.evaluate("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page2.wait_for_timeout(300)
        await page2.evaluate("() => __Q.openPane('styles')")
        await page2.wait_for_timeout(300)
        r = await page2.evaluate("() => ({ sw: document.documentElement.scrollWidth, cw: document.documentElement.clientWidth, n: document.querySelectorAll('.st-item').length })")
        ok(r['sw'] <= r['cw'] and r['n'] >= 12, 'on a phone the Styles pane fits', r)
        await page2.screenshot(path='shots/sty1-phone.png')
        for e in errors2: print('ERR(phone)', e)
        await b2.close()

        await page.screenshot(path='shots/sty1.png')
        errs = [e for e in errors if 'favicon' not in e[1]]
        ok(not errs, 'no errors in the console', errs[:5])
        await b.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')

asyncio.run(main())
