"""Section breaks. Layout › Breaks (Next Page, Continuous, Even Page, Odd Page) and what a section
has of its own: orientation, size and margins (the Layout menus and Page Setup's Apply to: this
section, this point forward, the whole document), its text column on screen, sheets of other
sizes centred, headers and footers linked to the section before or its own (Link to Previous,
Header -Section 2- / Same as Previous), Different First Page per section, page numbers that go on
or start again in a format of their own (Page Number Format), undo, copy and paste, deleting a
break, saved and opened again; the status bar; PDF (vector, pictures, print preview) with pages
of each size; Word files out and back (sectPr per section, LibreOffice), sections made by Word
(python-docx), columns still columns; the web page (named pages), Markdown and text; read-only;
phone width."""
import asyncio, sys, json, base64, io, zipfile, re, subprocess, os, tempfile
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)[:900]))
    if not cond: fails.append(what)

PARA = 'The office will remain closed on Sunday for the national holiday. All staff are requested to note this and plan their work accordingly. '
HTML = ('<h1>Office order</h1>' + ''.join(f'<p>{PARA * 2}</p>' for _ in range(3)) +
        '<p id="annex">Annex A: staff list</p><table style="width:100%"><tbody>' +
        ''.join(f'<tr><td>{i}</td><td>Name {i}</td><td>Section</td><td>Phone</td></tr>' for i in range(1, 8)) +
        '</tbody></table><p id="end">End of annex.</p><p id="back">Back to the order.</p>')
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""
OPEN_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); r.title = 'Imported'; await __Q.openImported(r); return true; }"""
GET_PDF = """async () => { await document.fonts.ready; const blob = await __Q.buildVectorPdf({ title: 'Sections' }); const u = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < u.length; i += 0x8000) s += String.fromCharCode.apply(null, u.subarray(i, i + 0x8000)); return btoa(s); }"""
CARET = """([id, end]) => { const p = typeof id === 'number' ? document.querySelectorAll('#editor > p')[id] : document.getElementById(id); const t = end ? p.lastChild : p.firstChild; const r = document.createRange(); r.setStart(t, end ? (t.length || 0) : 0); r.collapse(true); __Q.setRange(r); document.getElementById('editor').focus(); }"""
STATE = """() => ({ pages: __Q.S.pages, secOf: __Q.Pages.list.map(e => e.s), secs: __Q.Sections.list().map(s => [s.type, s.pg.orient, s.first, s.pn]),
  sheets: [...document.querySelectorAll('#sheets .sheet')].map(s => [Math.round(parseFloat(s.style.left)), Math.round(parseFloat(s.style.width)), Math.round(parseFloat(s.style.height))]),
  hdr: [...document.querySelectorAll('#sheets .sheet')].map(s => { const x = s.querySelector('.hfx.hdr'); return x ? x.textContent : ''; }),
  ftr: [...document.querySelectorAll('#sheets .sheet')].map(s => { const x = s.querySelector('.hfx.ftr'); return x ? x.textContent : ''; }),
  status: document.getElementById('sbPage').textContent })"""
WIDTH = "id => Math.round(document.getElementById(id).getBoundingClientRect().width / __Q.S.zoom)"

def pdf_sizes(path):
    out = subprocess.run(['pdfinfo', '-f', '1', '-l', '20', path], capture_output=True, text=True).stdout
    return [tuple(round(float(x)) for x in m) for m in re.findall(r'Page\s+\d+ size:\s+([\d.]+) x ([\d.]+)', out)]

def word_fixture():
    """three sections as Word writes them: portrait with a header, a landscape annex on a new page
    with a header of its own and page numbers from i, then a continuous section linked to it"""
    import docx
    from docx.enum.section import WD_SECTION, WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    s1 = d.sections[0]
    s1.header.paragraphs[0].text = 'Ministry of Planning'
    for i in range(3): d.add_paragraph(f'Order paragraph {i + 1}.')
    s2 = d.add_section(WD_SECTION.NEW_PAGE)
    s2.orientation = WD_ORIENT.LANDSCAPE
    s2.page_width, s2.page_height = s1.page_height, s1.page_width
    s2.header.is_linked_to_previous = False
    s2.header.paragraphs[0].text = 'Annex header'
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:start'), '1'); pg.set(qn('w:fmt'), 'lowerRoman')
    s2._sectPr.append(pg)
    d.add_paragraph('Annex paragraph.')
    s3 = d.add_section(WD_SECTION.CONTINUOUS)
    s3.orientation = WD_ORIENT.LANDSCAPE
    s3.page_width, s3.page_height = s2.page_width, s2.page_height
    s3.left_margin = s3.right_margin = docx.shared.Inches(2)
    # (python-docx copies the section before: this one goes on numbering)
    for el in s3._sectPr.findall(qn('w:pgNumType')): s3._sectPr.remove(el)
    d.add_paragraph('A narrower part of the annex.')
    bio = io.BytesIO(); d.save(bio)
    return base64.b64encode(bio.getvalue()).decode()

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1366, height=900, mock_caps=True, gfonts=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(500)
        s0 = await st(STATE)

        print('1. Layout › Breaks › Section Breaks')
        await st(CARET, ['annex', False])
        await page.click('#tab-layout')
        await page.locator('.rb:visible:has(.cap:text-is("Breaks"))').click()
        await page.wait_for_timeout(200)
        items = await st("() => [...document.querySelectorAll('.pop .mi .lb-sample')].map(e => e.firstChild.textContent)")
        ok(items == ['Next Page', 'Continuous', 'Even Page', 'Odd Page'], 'the Breaks menu: Next Page, Continuous, Even Page, Odd Page', items)
        await page.locator('.pop .mi', has_text='Next Page').click()
        await page.wait_for_timeout(400)
        r = await st("() => { const hr = document.querySelector('#editor > hr.sb'); return [hr && hr.dataset.sb, hr && JSON.parse(hr.dataset.pg).orient, hr && !!hr.dataset.sid, hr && hr.nextElementSibling.id, __Q.S.pages] }")
        ok(r == ['next', 'portrait', True, 'annex', 2], 'a Next Page break before the annex: it starts page 2, with the page setup of the section it was in', r)
        s1 = await st(STATE)
        ok(s1['secOf'] == [0, 1] and s1['status'] == 'Page 2 of 2 · Section 2', 'page 2 is section 2; the status bar says so', [s1['secOf'], s1['status']])
        await st("() => __Q.CMD.marks()")
        await page.wait_for_timeout(200)
        lab = await st("() => getComputedStyle(document.querySelector('#editor hr.sb'), '::after').content")
        ok(lab == '"Section Break (Next Page)"', 'with formatting marks it shows as Section Break (Next Page)', lab)
        await st("() => __Q.CMD.marks()")
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)
        ok(await st("() => [document.querySelectorAll('#editor hr.sb').length, __Q.S.pages]") == [0, s0['pages']], 'Undo takes the break away')
        await st("() => __Q.redo()"); await page.wait_for_timeout(300)
        ok(await st("() => document.querySelectorAll('#editor hr.sb').length") == 1, 'Redo puts it back')

        print('2. Orientation, size and margins of a section')
        await st(CARET, ['annex', False])
        await page.locator('.rb:visible:has(.cap:text-is("Orientation"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .mi', has_text='Landscape').click()
        await page.wait_for_timeout(400)
        s2 = await st(STATE)
        ok(s2['sheets'][0][1:] == [794, 1123] and s2['sheets'][1][1:] == [1123, 794], 'Layout › Orientation › Landscape: page 2 (section 2) turns, page 1 stays portrait', s2['sheets'])
        ok(s2['sheets'][0][0] == round((1123 - 794) / 2) and s2['sheets'][1][0] == 0, 'pages of different widths are centred in the column of pages', s2['sheets'])
        w = await st("() => [document.querySelector('#editor > p').getBoundingClientRect().width, document.getElementById('annex').getBoundingClientRect().width, document.querySelector('#editor table').getBoundingClientRect().width].map(x => Math.round(x / __Q.S.zoom))")
        ok(abs(w[0] - 602) <= 1 and abs(w[1] - 931) <= 1 and abs(w[2] - 931) <= 2, 'each section\'s text has its own column: 6.27" portrait, 9.69" landscape; the 100% table spans it', w)
        x = await st("() => { const a = document.querySelector('#editor > p').getBoundingClientRect(), b = document.getElementById('annex').getBoundingClientRect(), sh = [...document.querySelectorAll('#sheets .sheet')].map(s => s.getBoundingClientRect()), z = __Q.S.zoom; return [Math.round((a.left - sh[0].left) / z), Math.round((b.left - sh[1].left) / z)]; }")
        ok(x == [96, 96], 'and starts at its page\'s left margin', x)
        saved = await st("() => __Q.cleanHTML()")
        ok('data-sx' not in saved and '--sx' not in saved, 'the column is the layout\'s: nothing of it is saved')
        await page.locator('.rb:visible:has(.cap:text-is("Orientation"))').click()
        await page.wait_for_timeout(150)
        menu = await st("() => [...document.querySelectorAll('.pop .mi')].map(e => [e.textContent, e.getAttribute('aria-checked')])")
        await page.keyboard.press('Escape')
        ok(['Landscape', 'true'] in menu, 'the Orientation menu shows the caret\'s section', menu)
        await page.locator('.rb:visible:has(.cap:text-is("Margins"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .mi', has_text='Narrow').click()
        await page.wait_for_timeout(400)
        r = await st("() => __Q.Sections.list().map(s => [s.pg.ml, s.pg.mt])")
        ok(r == [[96, 96], [48, 48]], 'Margins › Narrow: only the section with the caret', r)
        ok(abs(await st(WIDTH, 'annex') - (1123 - 96)) <= 1, 'its column widens with them')
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)
        ok(await st("() => __Q.Sections.list()[1].pg.ml") == 96, 'Undo puts the margins back')

        print('3. Continuous, Even Page and Odd Page')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Breaks', html: h }))", '<p>First part.</p><p id="b">Second part, set narrower.</p><p>Third.</p>')
        await page.wait_for_timeout(300)
        await st(CARET, ['b', False])
        await st("() => __Q.Sections.insert('cont')")
        await st(CARET, ['b', False])
        await st("() => __Q.setPage({ ml: 192, mr: 192 })")
        await page.wait_for_timeout(300)
        r = await st("() => [__Q.S.pages, Math.round(document.querySelector('#editor > p').getBoundingClientRect().width / __Q.S.zoom), Math.round(document.getElementById('b').getBoundingClientRect().width / __Q.S.zoom)]")
        ok(r == [1, 602, 410], 'a Continuous break: the new section goes on on the same page, in a narrower column', r)
        for kind, want in [('even', 2), ('odd', 3)]:
            await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Breaks', html: h }))", '<p>Page one.</p><p id="b">The next section.</p>')
            await page.wait_for_timeout(300)
            await st(CARET, ['b', False])
            await st(f"() => __Q.Sections.insert('{kind}')")
            await page.wait_for_timeout(300)
            r = await st("() => { const rr = document.createRange(); rr.selectNodeContents(document.getElementById('b').firstChild); return [__Q.S.pages, __Q.Pages.list.map(e => e.s), __Q.Pages.at((rr.getBoundingClientRect().top - document.getElementById('paper').getBoundingClientRect().top) / __Q.S.zoom + 2)]; }")
            ok(r[0] == want and r[2] == want - 1, f'{kind.capitalize()} Page: the section starts on page {want}' + (' (page 2 left blank)' if kind == 'odd' else ''), r)

        print('4. Headers and footers of a section')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(400)
        await st(CARET, ['annex', False])
        await st("() => __Q.Sections.insert('next')")
        await page.wait_for_timeout(300)
        await st("() => __Q.HF.edit('header', 0)")
        await page.keyboard.type('Ministry of Planning')
        await st("() => __Q.HF.close()")
        await page.wait_for_timeout(300)
        ok((await st(STATE))['hdr'] == ['Ministry of Planning', 'Ministry of Planning'], 'section 2 shows the header of section 1 (linked, as Word starts it)')
        await st("() => __Q.HF.edit('header', 1)")
        await page.wait_for_timeout(300)
        tags = await st("() => [...document.querySelectorAll('.hf-band .hf-tag')].map(x => x.textContent)")
        ok(tags == ['Header -Section 2-', 'Same as Previous'], 'its header is tagged Header -Section 2- and Same as Previous', tags)
        btn = page.locator('.rb:visible:has(.cap:text-is("Link to Previous"))')
        ok(await btn.get_attribute('aria-pressed') == 'true', 'Link to Previous is on')
        await btn.click()
        await page.wait_for_timeout(300)
        ok(await st("() => [...document.querySelectorAll('.hf-band .hf-tag')].map(x => x.textContent)") == ['Header -Section 2-'] and await btn.get_attribute('aria-pressed') == 'false', 'turned off: a header of its own (a copy to change)')
        await page.keyboard.press('Control+a')
        await page.keyboard.type('Annex A')
        await st("() => __Q.HF.close()")
        await page.wait_for_timeout(300)
        ok((await st(STATE))['hdr'] == ['Ministry of Planning', 'Annex A'], 'changed in section 2 only', (await st(STATE))['hdr'])
        # Different First Page for section 2, page numbers in the footer
        await st(CARET, ['annex', False])
        await st("() => __Q.HF.pageNumbers('footer', 'center')")
        await page.wait_for_timeout(300)
        ok((await st(STATE))['ftr'] == ['1', '2'], 'Page Number: numbers on both sections\' pages (the footer is linked)', (await st(STATE))['ftr'])
        await st("() => { __Q.HF.edit('footer', 1); __Q.CMD.hfFirst(); }")
        await page.wait_for_timeout(300)
        r = await st("() => [__Q.Sections.list().map(s => s.first), __Q.HF.variant(0), __Q.HF.variant(1)]")
        ok(r == [[False, True], 'def', 'first'], 'Different First Page for section 2 only: its first page has its own footer', r)
        await st("() => __Q.HF.close()")
        await page.wait_for_timeout(300)
        ok((await st(STATE))['ftr'] == ['1', ''], 'empty, as Word starts it', (await st(STATE))['ftr'])
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)

        print('5. Page numbers of a section')
        await st(CARET, ['annex', False])
        await st("() => __Q.Dialogs.pageNumberFormat()")
        await page.wait_for_timeout(200)
        ok(await st("() => document.querySelector('.dialog h2').textContent") == 'Page Number Format', 'Page Number Format opens')
        await page.select_option('#pnFmt', 'i')
        await page.select_option('#pnMode', 'start')
        await page.fill('#pnStart', '1')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        s5 = await st(STATE)
        ok(s5['ftr'] == ['1', 'i'] and s5['secs'][1][3] == {'start': 1, 'fmt': 'i'}, 'section 2 numbered from i (roman), section 1 from 1', s5)
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)
        ok((await st(STATE))['ftr'] == ['1', '2'], 'Undo: numbers go on again')
        await st("() => __Q.redo()"); await page.wait_for_timeout(300)

        print('6. Page Setup: this section, this point forward, the whole document')
        await st(CARET, ['annex', False])
        await st("() => __Q.Dialogs.pageSetup()")
        await page.wait_for_timeout(200)
        opts = await st("() => [...document.querySelectorAll('#psApply option')].map(o => o.textContent)")
        ok(opts == ['This section', 'This point forward', 'Whole document'], 'Apply to: This section, This point forward, Whole document', opts)
        await page.select_option('#psOrient', 'landscape')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        r = await st("() => __Q.Sections.list().map(s => [s.pg.orient, Math.round(s.g.w)])")
        ok(r == [['portrait', 794], ['landscape', 1123]], 'This section: landscape (A4 kept exact)', r)
        await st(CARET, [2, False])
        await st("() => __Q.Dialogs.pageSetup()")
        await page.wait_for_timeout(200)
        await page.select_option('#psApply', 'fwd')
        await page.select_option('#psSize', 'Letter')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        r = await st("() => [document.querySelectorAll('#editor > hr.sb').length, __Q.Sections.list().map(s => s.pg.size), __Q.Sections.indexOf(document.querySelectorAll('#editor > p')[2])]")
        ok(r == [2, ['A4', 'Letter', 'Letter'], 1], 'This point forward: a new section from the caret, it and those after it on Letter paper', r)
        await st("() => __Q.Dialogs.pageSetup()")
        await page.wait_for_timeout(200)
        await page.select_option('#psApply', 'all')
        await page.fill('#psTop', '0.5"')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(400)
        r = await st("() => __Q.Sections.list().map(s => s.pg.mt)")
        ok(r == [48, 48, 48], 'Whole document: every section', r)
        await st("() => { __Q.undo(); __Q.undo(); }"); await page.wait_for_timeout(400)
        ok(await st("() => document.querySelectorAll('#editor > hr.sb').length") == 1, 'Undo twice: back to two sections')

        print('7. Copy, paste, delete; saved and opened again')
        sid = await st("() => document.querySelector('#editor > hr.sb').dataset.sid")
        await st("() => { const hr = document.querySelector('#editor > hr.sb'); const r = document.createRange(); r.setStartBefore(hr.previousElementSibling); r.setEndAfter(hr); __Q.setRange(r); }")
        await page.keyboard.press('Control+c')
        await st(CARET, ['back', True])
        await page.keyboard.press('Control+v')
        await page.wait_for_timeout(400)
        r = await st("() => { const hs = [...document.querySelectorAll('#editor > hr.sb')]; return [hs.length, new Set(hs.map(h => h.dataset.sid)).size, Object.keys(__Q.S.doc.hf.sects || {}).length]; }")
        ok(r == [2, 2, 2], 'a copied break pasted: a section of its own (new id; its own header copied)', r)
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)
        # Backspace at the start of section 2: the break goes, the sections join
        await st(CARET, ['annex', False])
        await page.keyboard.press('Backspace')
        await page.wait_for_timeout(400)
        r = await st("() => [document.querySelectorAll('#editor > hr.sb').length, Math.round(parseFloat(document.getElementById('editor').style.width)), Math.round(document.getElementById('annex').getBoundingClientRect().width / __Q.S.zoom), [...document.querySelectorAll('#sheets .sheet')].map(s => Math.round(parseFloat(s.style.width)))]")
        ok(r[0] == 0 and r[1] == 794 and abs(r[2] - 602) <= 1 and set(r[3]) == {794}, 'Backspace at the start of a section takes its break away: one section, all portrait again', r)
        snap = await st("() => { const s = __Q.snapshotDoc(); return [Object.keys((s.hf && s.hf.sects) || {}).length, Object.keys(__Q.S.doc.hf.sects || {}).length]; }")
        ok(snap == [0, 1], 'the saved copy keeps no header for a section that is gone (the open one does, for Undo)', snap)
        await st("() => __Q.undo()"); await page.wait_for_timeout(300)
        r = await st("""async () => { await __Q.Saver.flush(); __Q.Saver.mark(); await __Q.Saver.flush(); const back = await __Q.CloudStore.load(__Q.S.doc.id); __Q.openDoc(back); await new Promise(r => setTimeout(r, 500)); return [document.querySelectorAll('#editor > hr.sb').length, __Q.Sections.list().map(s => [s.pg.orient, s.pn])]; }""")
        ok(r == [1, [['portrait', None], ['landscape', {'start': 1, 'fmt': 'i'}]]], 'saved to the account and opened again: the section, its page setup and numbers', r)
        s7 = await st(STATE)
        ok(s7['hdr'] == ['Ministry of Planning', 'Annex A'] and s7['ftr'] == ['1', 'i'], 'and its own header', [s7['hdr'], s7['ftr']])

        print('8. PDF and print')
        open('shots/sec1.pdf', 'wb').write(base64.b64decode(await st(GET_PDF)))
        sizes = pdf_sizes('shots/sec1.pdf')
        ok(sizes == [(595, 842), (842, 595)], 'the vector PDF: an A4 page, then an A4 landscape one', sizes)
        t2 = subprocess.run(['pdftotext', '-layout', '-f', '2', '-l', '2', 'shots/sec1.pdf', '-'], capture_output=True, text=True).stdout
        ok('Annex A' in t2 and 'Name 7' in t2 and re.search(r'^\s*i\s*$', t2, re.M), 'page 2: its header, the table, page number i', t2[:400])
        r = await st("async () => { const pages = await __Q.renderPagesForPdf({ scale: 1 }); return pages.map(p => [Math.round(p.wPt), Math.round(p.hPt)]); }")
        ok(r == [[595, 842], [842, 595]], 'the picture PDF too', r)
        r = await st("async () => { const P = await __Q.collectPaint(); const c = __Q.paintPage(P, 1, 1); return [c.width, c.height]; }")
        ok(r == [1123, 794], 'the page painter (print preview) paints page 2 landscape', r)

        print('9. Word')
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        doc = z.read('word/document.xml').decode()
        sects = re.findall(r'<w:sectPr>.*?</w:sectPr>', doc)
        ok(len(sects) == 2 and 'w:orient="landscape"' in sects[1] and 'w:orient' not in sects[0], 'two sections: portrait, then landscape', [x[:200] for x in sects])
        inpara = re.search(r'<w:pPr>(?:(?!</w:pPr>).)*<w:sectPr>', doc) is not None
        empty = re.search(r'<w:p><w:pPr><w:sectPr>(?:(?!</w:sectPr>).)*</w:sectPr></w:pPr></w:p>', doc) is not None
        ok(inpara and not empty and '<w:type w:val="nextPage"/>' in sects[1], 'the first section\'s settings in its last paragraph (no empty one added); the second starts on a new page', [inpara, empty])
        ok('<w:pgNumType w:fmt="lowerRoman" w:start="1"/>' in sects[1], 'page numbers from i in section 2', sects[1][:600])
        ok(sects[0].count('headerReference') == 1 and sects[1].count('headerReference') == 1 and sects[1].count('footerReference') == 0, 'section 2 has a header of its own; its footer is section 1\'s (linked)')
        await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(800)
        s9 = await st(STATE)
        ok(s9['secs'] == [['next', 'portrait', False, None], ['next', 'landscape', False, {'start': 1, 'fmt': 'i'}]] and s9['hdr'] == ['Ministry of Planning', 'Annex A'] and s9['ftr'] == ['1', 'i'], 'opened again: the same sections, headers and numbers', s9)
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, 'sec.docx'); open(src, 'wb').write(base64.b64decode(b64))
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', td, src], capture_output=True, timeout=120)
            lo = os.path.join(td, 'sec.pdf')
            sizes = pdf_sizes(lo) if os.path.exists(lo) else []
            ok(sizes == [(595, 842), (842, 595)], 'LibreOffice: an A4 page and a landscape one', sizes)
            t = subprocess.run(['pdftotext', '-layout', lo, '-'], capture_output=True, text=True).stdout if os.path.exists(lo) else ''
            ok('Ministry of Planning' in t and 'Annex A' in t and re.search(r'^\s*i\s*$', t, re.M), 'with each section\'s header and its page numbers', t[-500:])

        print('10. Sections made by Word')
        await st(OPEN_DOCX, word_fixture())
        await page.wait_for_timeout(900)
        s10 = await st(STATE)
        r = await st("() => __Q.Sections.list().map(s => [s.type, s.pg.orient, s.pg.ml, s.pn])")
        ok(r == [['next', 'portrait', 120, None], ['next', 'landscape', 120, {'start': 1, 'fmt': 'i'}], ['cont', 'landscape', 192, None]], 'three sections: portrait; landscape on a new page, numbered from i; continuous with wider margins', r)
        ok(s10['hdr'][:2] == ['Ministry of Planning', 'Annex header'] and s10['pages'] == 2, 'each with the header Word gave it (the third linked to the second), on two pages', s10)
        w = await st("() => [...document.querySelectorAll('#editor > p')].map(p => Math.round(p.getBoundingClientRect().width / __Q.S.zoom))")
        cw = await st("() => __Q.Sections.list().map(s => Math.round(s.g.cw))")
        ok(abs(w[-1] - cw[2]) <= 1 and abs(w[-2] - cw[1]) <= 1 and abs(w[0] - cw[0]) <= 1 and cw[2] < cw[1], 'the continuous section\'s narrower column on the same page', [w, cw])
        # columns are still columns
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Cols', html: h }))", '<p>Before.</p><div class="cols" style="--cols:2"><p>One.</p><p>Two.</p></div><p>After.</p>')
        await page.wait_for_timeout(300)
        await st(OPEN_DOCX, await st(GET_DOCX))
        await page.wait_for_timeout(600)
        r = await st("() => [document.querySelectorAll('#editor > div.cols').length, document.querySelectorAll('#editor hr.sb').length, document.querySelector('#editor > div.cols') && document.querySelector('#editor > div.cols').style.getPropertyValue('--cols')]")
        ok(r == [1, 0, '2'], 'two columns in a part of the page, out to Word and back: still a part in columns (no section break)', r)
        # with a header (Word's parts in columns name it again)
        await st("() => { __Q.change(() => { __Q.S.doc.hf.header.def = '<p>Letterhead</p>'; }, { full: true }); __Q.HF.refresh(); }")
        await page.wait_for_timeout(300)
        await st(OPEN_DOCX, await st(GET_DOCX))
        await page.wait_for_timeout(600)
        r = await st("() => [document.querySelectorAll('#editor > div.cols').length, document.querySelectorAll('#editor hr.sb').length, __Q.S.doc.hf.header.def.includes('Letterhead')]")
        ok(r == [1, 0, True], 'and with a header too', r)

        print('11. Web page, Markdown, text')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(300)
        await st(CARET, ['annex', False])
        await st("() => __Q.Sections.insert('next')")
        await st(CARET, ['annex', False])
        await st("() => __Q.setPage({ orient: 'landscape' })")
        await page.wait_for_timeout(300)
        html = await st("() => __Q.toStandaloneHtml()")
        ok('@page s1{size:297mm 210mm' in html and 'data-sect="1"' in html and 'hr.sb:not([data-sb="cont"]){break-after:page' in html, 'the web page prints section 2 on landscape pages (a named page) from a new page', re.findall(r'@page[^}]*}', html))
        md = await st("() => __Q.toMarkdown()"); tx = await st("() => __Q.toPlainText()")
        ok(not re.search(r'^---$', md, re.M) and 'Annex A' in md and '\f' in tx, 'Markdown: no rule for it; plain text: a page break', md[:200])

        print('12. Read-only, phone')
        await st("() => { __Q.S.readOnly = true; }")
        n0 = await st("() => document.querySelectorAll('#editor hr.sb').length")
        await st("() => __Q.Sections.insert('next')")
        ok(await st("() => document.querySelectorAll('#editor hr.sb').length") == n0, 'read-only: no section break')
        await st("() => { __Q.S.readOnly = false; }")
        errs = [e for e in errors if 'favicon' not in e[1]]
        ok(not errs, 'no errors in the console', errs[:5])
        await b.close()

        b2, page2, errors2 = await new_page(p, width=400, height=860, mock_caps=True, gfonts=True)
        await load(page2)
        await page2.evaluate("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page2.wait_for_timeout(400)
        await page2.click('#tab-layout')
        await page2.locator('.rb:visible:has(.cap:text-is("Breaks"))').click()
        await page2.wait_for_timeout(200)
        r = await page2.evaluate("() => { const m = document.querySelector('.pop'); const r = m.getBoundingClientRect(); return [r.left >= 0, r.right <= innerWidth, document.documentElement.scrollWidth <= innerWidth]; }")
        ok(r == [True, True, True], 'phone: the Breaks menu fits the screen', r)
        await page2.screenshot(path='shots/sec1-phone.png')
        for e in errors2: print('ERR(phone)', e)
        await b2.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')

asyncio.run(main())
