"""Headers and footers (letterheads): the model (old left/centre/right texts brought over);
Insert › Header letterheads (first page) with a logo and a table, the placeholder selected to type
over; editing in place (the page's text dims; Esc, Close or a double-click in the text puts it
back as one undo step); a header taller than the top margin pushes the page's text down; page
numbers (Bangla digits, Page X of Y) on every page but a different first page; Different Odd &
Even Pages; Previous / Next; read-only; the Word file (header and footer parts per variant, fields,
pictures with their own relationships, the watermark in every header, evenAndOddHeaders) out and
back in, a Word-made file with complex fields, a floating logo, LibreOffice's rendering of ours;
the vector PDF and the HTML page."""
import asyncio, sys, json, base64, io, zipfile, re, subprocess, os, tempfile
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)[:600]))
    if not cond: fails.append(what)

PARA = '<p>' + 'The office will remain closed on Sunday for the national holiday. All staff are requested to note this and plan their work accordingly. ' * 4 + '</p>'
HTML = '<p>Ref: ABC/2026/112</p><p>Date: 19 October 2026</p>' + PARA * 16
SHEETS = "() => [...document.querySelectorAll('#sheets .sheet')].map(s => [...s.querySelectorAll('.hfx')].map(x => (x.classList.contains('hdr') ? 'H:' : 'F:') + x.innerText.replace(/\\s+/g, ' ').trim()))"
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""
READ_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); return r; }"""
OPEN_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); r.title = 'Imported'; await __Q.openImported(r); return true; }"""
PNG = 'iVBORw0KGgoAAAANSUhEUgAAAAgAAAAICAIAAABLbSncAAAAGElEQVR4nGP4z8DwnwELYMQqMUglsDsJAJkQB/kYyQq1AAAAAElFTkSuQmCC'

def word_made():
    """a letter as Word makes it: a first-page letterhead (a table with the name), a header on the
    other pages, a footer with a PAGE field (fldChar), even pages different"""
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, Inches
    d = docx.Document()
    s = d.sections[0]
    s.different_first_page_header_footer = True
    d.settings.odd_and_even_pages_header_footer = True
    fh = s.first_page_header
    t = fh.add_table(1, 2, Inches(6))
    t.cell(0, 0).paragraphs[0].add_run('Bangladesh Bank').bold = True
    t.cell(0, 1).paragraphs[0].add_run('Head Office, Motijheel')
    fh.paragraphs[0].text = ''
    s.header.paragraphs[0].text = 'Bangladesh Bank\tCircular No. 7\tcontinued'
    s.even_page_header.paragraphs[0].text = 'Even page header text'
    def page_field(par, prefix):
        par.add_run(prefix)
        r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
        r = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE   \\* MERGEFORMAT '; r._r.append(it)
        r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'separate'); r._r.append(fc)
        r = par.add_run('3'); r.font.size = Pt(9)
        r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); r._r.append(fc)
    page_field(s.footer.paragraphs[0], 'Page ')
    page_field(s.even_page_footer.paragraphs[0], 'Even page ')
    for i in range(90): d.add_paragraph('Paragraph %d of the circular, long enough to fill the pages of the letter with ordinary text.' % (i + 1))
    bio = io.BytesIO(); d.save(bio); return base64.b64encode(bio.getvalue()).decode()

def pdf_pages_text(path):
    out = []
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', path], capture_output=True, text=True).stdout).group(1))
    for i in range(1, n + 1):
        out.append(subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', path, '-'], capture_output=True, text=True).stdout)
    return out

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1366, height=900, mock_caps=True, gfonts=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)

        print('1. The model; headers from before')
        r = await st("""() => { const d = __Q.newDocModel({ title: 'Old', header: { l: 'Folio', c: '', r: 'Page {PAGE} of {PAGES}' }, footer: { l: '', c: '{TITLE}', r: '' } }); return { hd: d.hf.header.def, fd: d.hf.footer.def, legacy: [d.header, d.footer], first: d.hf.header.first, even: d.evenDiff }; }""")
        ok('hf-grid' in r['hd'] and 'data-f="PAGE"' in r['hd'] and 'data-f="PAGES"' in r['hd'] and 'Folio' in r['hd'], 'left and right texts become a borderless three-cell table with fields', r['hd'])
        ok(r['fd'].startswith('<p style="text-align:center') and 'data-f="TITLE"' in r['fd'], 'a lone centre text becomes a centred paragraph', r['fd'])
        ok(r['legacy'] == [{'l': '', 'c': '', 'r': ''}] * 2 and r['first'] == '' and r['even'] is False, 'the old fields are emptied; no first-page or even header', r)
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Old', html: h, header: { l: 'Folio', c: '', r: 'Page {PAGE} of {PAGES}' } }))", HTML)
        await page.wait_for_timeout(500)
        sh = await st(SHEETS)
        n = len(sh)
        ok(n >= 3 and sh[0] == ['H:Folio Page 1 of %d' % n] and sh[1] == ['H:Folio Page 2 of %d' % n], 'each page shows the header with its own page number', sh)
        # code from before that sets the texts directly still works
        await st("() => { const d = __Q.S.doc; d.hf.header.def = ''; d.header = { l: '', c: 'Set the old way', r: '' }; __Q.renderSheets(); }")
        await page.wait_for_timeout(300)
        ok((await st(SHEETS))[0] == ['H:Set the old way'], 'a header set the old way is brought over when the pages are drawn', await st(SHEETS))

        print('2. A letterhead from Insert › Header')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(400)
        pages0 = await st("() => __Q.S.pages")
        await st("() => __Q.UI.selectTab('insert')")
        await page.locator('.rb:visible:has(.cap:text-is("Header"))').first.click()
        await page.wait_for_timeout(200)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(m => m.textContent.trim())")
        ok(any('logo on the left' in i for i in items) and any('Bangla and English' in i for i in items) and any('Remove Header' in i for i in items), 'the Header menu has Edit Header, the letterheads and Remove Header', items)
        await page.locator('.pop .mi', has_text='logo on the left').click()
        await page.wait_for_timeout(400)
        r = await st("() => ({ editing: __Q.HF.editing && [__Q.HF.editing.kind, __Q.HF.editing.variant, __Q.HF.editing.p], first: __Q.S.doc.firstDiff, tab: __Q.S.tab, sel: getSelection().toString(), isEditor: __Q.S.editor === (__Q.HF.editing && __Q.HF.editing.el), bodyCE: document.getElementById('editor').contentEditable, dim: document.getElementById('paper').classList.contains('hf-editing'), tag: document.querySelector('.hf-band .hf-tag') && document.querySelector('.hf-band .hf-tag').textContent })")
        ok(r['editing'] == ['header', 'first', 0] and r['first'] is True, 'a letterhead goes in the first page\'s header (Different First Page turned on)', r)
        ok(r['sel'] == '[Organization name]', 'its first placeholder is selected, to type over', r['sel'])
        ok(r['isEditor'] and r['bodyCE'] == 'false' and r['dim'] and r['tab'] == 'hf', 'the header is the editor; the page\'s text is dimmed and locked; the Header & Footer tab is open', r)
        ok(r['tag'] == 'First Page Header', 'the dashed line is tagged First Page Header', r['tag'])
        await page.keyboard.type('Ministry of Planning')
        await page.wait_for_timeout(300)
        r = await st("() => ({ pad: parseFloat(document.getElementById('editor').style.paddingTop), h: __Q.HF.heights['0:header:first'], g: __Q.geom() })")
        ok(abs(r['pad'] - (r['g']['hd'] + r['h'])) < 0.6 and r['pad'] > r['g']['mt'], 'the letterhead is taller than the top margin, so the text starts below it (header distance + its height)', r)
        # a click in the page's text doesn't take the caret out of the header
        box = await st("() => { const r = document.querySelectorAll('#editor > p')[3].getBoundingClientRect(); return { x: r.left + 40, y: r.top + 8 }; }")
        await page.mouse.click(box['x'], box['y'])
        await page.wait_for_timeout(150)
        r = await st("() => { const s = getSelection(); return !!(s.rangeCount && __Q.HF.editing && __Q.HF.editing.el.contains(s.anchorNode)); }")
        ok(r, 'a click in the dimmed text leaves the caret in the header')
        await page.keyboard.press('Escape')
        await page.wait_for_timeout(300)
        r = await st("() => ({ editing: !!__Q.HF.editing, tab: __Q.S.tab, ce: document.getElementById('editor').contentEditable, first: __Q.S.doc.hf.header.first, def: __Q.S.doc.hf.header.def, pages: __Q.S.pages, hctx: __Q.UI.ctx.hf, live: document.querySelectorAll('.hf-live, .hf-band').length })")
        ok(not r['editing'] and r['tab'] != 'hf' and not r['hctx'] and r['ce'] == 'true' and r['live'] == 0, 'Esc puts it back: the text is editable again and the tab goes', r)
        ok('Ministry of Planning' in r['first'] and '<img' in r['first'] and 'hf-grid' in r['first'] and r['def'] == '', 'the first page\'s header keeps the logo, the table and the typed name; the other pages have none', r['first'][:200])
        r2 = await st("() => { const s = document.querySelector('#sheets .sheet').getBoundingClientRect(), p = document.querySelector('#editor > p').getBoundingClientRect(); return [(p.top - s.top) / __Q.S.zoom, parseFloat(document.getElementById('editor').style.paddingTop)]; }")
        ok(abs(r2[0] - r2[1]) < 1, 'the first line of text is below the letterhead', r2)
        sh = await st(SHEETS)
        ok(len(sh[0]) == 1 and 'Ministry of Planning' in sh[0][0] and sh[1] == [], 'page 1 shows the letterhead; page 2 no header', sh[:2])
        r = await st("() => { const im = document.querySelector('#sheets .sheet .hfx img'); const td = document.querySelector('#sheets .sheet .hfx td'); return { img: !!im && im.complete && im.naturalWidth > 0, border: td && getComputedStyle(td).borderTopColor }; }")
        ok(r['img'] and r['border'] in ('rgba(0, 0, 0, 0)', 'transparent'), 'the logo shows on the page; the table\'s gridlines don\'t', r)
        # undo / redo: one step
        await st("() => __Q.undo()")
        await page.wait_for_timeout(300)
        r = await st("() => ({ first: __Q.S.doc.hf.header.first, sh: [...document.querySelectorAll('#sheets .sheet')][0].querySelectorAll('.hfx').length })")
        ok(r['first'] == '' and r['sh'] == 0, 'Undo takes the header\'s editing back in one step', r)
        await st("() => __Q.redo()")
        await page.wait_for_timeout(300)
        ok('Ministry of Planning' in await st("() => __Q.S.doc.hf.header.first"), 'Redo brings it back')

        print('3. Page numbers, the other pages\' header, Previous / Next')
        await st("() => __Q.UI.selectTab('insert')")
        await page.locator('.rb:visible:has(.cap:text-is("Page Number"))').first.click()
        await page.wait_for_timeout(200)
        await page.locator('.pop .mi', has_text='Bottom center: পৃষ্ঠা ১').click()
        await page.wait_for_timeout(400)
        sh = await st(SHEETS)
        n = len(sh)
        ok(sh[0] == ['H:' + sh[0][0][2:]] and all(sh[i] == ['F:পৃষ্ঠা ' + '০১২৩৪৫৬৭৮৯'[i + 1]] for i in range(1, min(n, 9))), 'Bangla page numbers on every page but the first (Different First Page)', sh)
        r = await st("() => { const i0 = __Q.Hist.idx, h0 = __Q.S.doc.hf.footer.def; __Q.HF.edit('footer', 2); __Q.HF.close(); return [i0, __Q.Hist.idx, h0 === __Q.S.doc.hf.footer.def]; }")
        ok(r[0] == r[1] and r[2], 'opening page 3\'s footer and closing it changes nothing (its number is filled in per page)', r)
        # double-click in page 2's top margin: its header
        await st("() => document.querySelectorAll('#sheets .sheet')[1].scrollIntoView({ block: 'start' })")
        await page.wait_for_timeout(200)
        pt = await st("() => { const r = document.querySelectorAll('#sheets .sheet')[1].getBoundingClientRect(); return { x: r.left + r.width / 2, y: r.top + 40 }; }")
        await page.mouse.dblclick(pt['x'], pt['y'])
        await page.wait_for_timeout(300)
        r = await st("() => __Q.HF.editing && [__Q.HF.editing.kind, __Q.HF.editing.variant, __Q.HF.editing.p, document.querySelector('.hf-band .hf-tag').textContent]")
        ok(r == ['header', 'def', 1, 'Header'], 'a double-click in page 2\'s top margin edits the header of the other pages', r)
        await page.keyboard.type('Office order 12/2026')
        # Previous goes to the first page's header (a different one), Next comes back
        await st("() => __Q.CMD.hfPrev()")
        await page.wait_for_timeout(250)
        r = await st("() => __Q.HF.editing && [__Q.HF.editing.variant, __Q.HF.editing.p]")
        ok(r == ['first', 0], 'Previous: the first page\'s header', r)
        await st("() => __Q.CMD.hfNext()")
        await page.wait_for_timeout(250)
        r = await st("() => __Q.HF.editing && [__Q.HF.editing.variant, __Q.HF.editing.p, __Q.HF.editing.el.textContent]")
        ok(r and r[0] == 'def' and r[1] == 1 and 'Office order 12/2026' in r[2], 'Next: back to the other pages\' header, with what was typed', r)
        await st("() => __Q.CMD.hfToFooter()")
        await page.wait_for_timeout(250)
        r = await st("() => __Q.HF.editing && [__Q.HF.editing.kind, __Q.HF.editing.variant, __Q.HF.editing.el.textContent, getComputedStyle(__Q.HF.editing.el).transform !== 'none']")
        ok(r and r[0] == 'footer' and r[1] == 'def' and 'পৃষ্ঠা' in r[2], 'Go to Footer: the footer with the page number', r)
        # a double-click in the text closes
        await st("() => document.querySelectorAll('#editor > p')[6].scrollIntoView({ block: 'center' })")
        await page.wait_for_timeout(200)
        pt = await st("() => { const r = document.querySelectorAll('#editor > p')[6].getBoundingClientRect(); return { x: r.left + 50, y: r.top + 6 }; }")
        await page.mouse.dblclick(pt['x'], pt['y'])
        await page.wait_for_timeout(300)
        sh = await st(SHEETS)
        ok(not await st("() => !!__Q.HF.editing") and sh[1] == ['H:Office order 12/2026', 'F:পৃষ্ঠা ২'] and 'Ministry' in sh[0][0], 'a double-click in the text closes; page 2 shows its header and number', sh[:2])

        print('4. Different Odd & Even Pages')
        await st("() => __Q.HF.setEven(true)")
        await page.wait_for_timeout(300)
        r = await st("() => ({ even: __Q.S.doc.hf.header.even, ev: __Q.S.doc.evenDiff, f: __Q.S.doc.hf.footer.even })")
        ok(r['ev'] and 'Office order 12/2026' in r['even'] and 'data-f="PAGE"' in r['f'], 'even pages start with a copy of the header and footer', r)
        await st("() => __Q.HF.edit('header', 1)")
        await page.wait_for_timeout(250)
        r = await st("() => [__Q.HF.editing.variant, document.querySelector('.hf-band .hf-tag').textContent]")
        ok(r == ['even', 'Even Page Header'], 'page 2 is an even page', r)
        await st("() => { const el = __Q.HF.editing.el; const r = document.createRange(); r.selectNodeContents(el); getSelection().removeAllRanges(); getSelection().addRange(r); }")
        await page.keyboard.type('Even pages: Office order 12/2026')
        await st("() => __Q.HF.close()")
        await page.wait_for_timeout(300)
        sh = await st(SHEETS)
        ok(sh[1][0] == 'H:Even pages: Office order 12/2026' and sh[2][0] == 'H:Office order 12/2026', 'even and odd pages show their own headers', sh[:3])
        await st("() => __Q.HF.edit('header', 2)")
        await page.wait_for_timeout(200)
        ok(await st("() => document.querySelector('.hf-band .hf-tag').textContent") == 'Odd Page Header', 'the others are tagged Odd Page Header')
        await page.locator('.rb:visible:has(.cap:text-is("Close"))').first.click()
        await page.wait_for_timeout(250)
        ok(not await st("() => !!__Q.HF.editing"), 'Close on the Header & Footer tab closes')

        print('5. Read-only; the old dialog; Remove')
        await st("() => { __Q.S.readOnly = true; }")
        r = await st("() => __Q.HF.edit('header')")
        ok(r is False and not await st("() => !!__Q.HF.editing"), 'a read-only document\'s header can\'t be edited')
        await st("() => { __Q.S.readOnly = false; }")
        await st("() => __Q.Dialogs.headerFooter('footer')")
        await page.wait_for_timeout(200)
        ok(await st("() => __Q.HF.editing && __Q.HF.editing.kind") == 'footer', 'the Header & Footer command (old dialog) opens the footer on the page')
        await page.keyboard.press('Escape')
        await page.wait_for_timeout(200)

        print('6. The Word file')
        await st("() => __Q.setWatermark({ text: 'CONFIDENTIAL', color: '#c0c0c0', opacity: 0.4 })")
        await page.wait_for_timeout(200)
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        names = z.namelist()
        docx = z.read('word/document.xml').decode()
        sect = re.findall(r'<w:sectPr>.*?</w:sectPr>', docx)[-1]
        refs = re.findall(r'<w:(header|footer)Reference w:type="(\w+)" r:id="(\w+)"/>', sect)
        ok(sorted((k, t) for k, t, _ in refs) == [('footer', 'default'), ('footer', 'even'), ('header', 'default'), ('header', 'even'), ('header', 'first')], 'the section names the default, first-page and even headers and footers (the empty first-page footer left out)', refs)
        ok('<w:titlePg/>' in sect and '<w:evenAndOddHeaders/>' in z.read('word/settings.xml').decode(), 'Different First Page (titlePg) and Odd & Even (evenAndOddHeaders)')
        rels = z.read('word/_rels/document.xml.rels').decode()
        parts = {}
        for k, t, rid in refs:
            tgt = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid, rels).group(1)
            parts[(k, t)] = tgt
        first = z.read('word/' + parts[('header', 'first')]).decode()
        ok('<w:tbl>' in first and 'Ministry of Planning' in first and first.count('<w:drawing>') == 2 and '<w:pStyle w:val="Header"/>' in first, 'the letterhead part: the table, the name, the logo and the watermark, in the Header style', first[:300])
        frels = z.read('word/_rels/' + parts[('header', 'first')] + '.rels').decode()
        ok(re.search(r'Target="media/image\d+\.png"', frels) and 'rIdWm' in frels, 'its pictures are in its own relationships', frels)
        logo_t = re.search(r'Id="(rId\d+)"[^>]*Target="(media/image\d+\.png)"', frels)
        ok(logo_t and logo_t.group(2) in [ 'word/' + x.split('word/')[-1] for x in names] or ('word/' + logo_t.group(2)) in names, 'the logo file is in the package', logo_t and logo_t.group(2))
        dfl = z.read('word/' + parts[('header', 'default')]).decode()
        ok('Office order 12/2026' in dfl and 'name="Watermark"' in dfl and 'descr="CONFIDENTIAL"' in dfl, 'every header carries the watermark', dfl[:200])
        foot = z.read('word/' + parts[('footer', 'default')]).decode()
        ok('<w:fldSimple w:instr=" PAGE \\* MERGEFORMAT ">' in foot and '>১<' in foot and '<w:pStyle w:val="Footer"/>' in foot, 'the page number is a PAGE field (shown in Bangla digits), in the Footer style', foot[:400])
        ct = z.read('[Content_Types].xml').decode()
        ok(all(('/word/' + v) in ct for v in parts.values()), 'each part is in the content types')
        open('shots/hf1.docx', 'wb').write(base64.b64decode(b64))
        # back in
        res = await st(READ_DOCX, b64)
        s = res['settings']
        hf = s.get('hf') or {}
        ok(s.get('firstDiff') and s.get('evenDiff'), 'Different First Page and Odd & Even come back', [s.get('firstDiff'), s.get('evenDiff')])
        hfirst = hf.get('header', {}).get('first', '')
        ok('Ministry of Planning' in hfirst and '<table' in hfirst and re.search(r'<img[^>]+src="data:image/', hfirst), 'the letterhead comes back with its table and logo', hfirst[:300])
        ok('Watermark' not in json.dumps(hf) and s.get('watermark', {}).get('text') == 'CONFIDENTIAL', 'the watermark comes back as the watermark, not a picture in the header', s.get('watermark'))
        fdef = hf.get('footer', {}).get('def', '')
        ok(re.search(r'<span class="hff" data-f="PAGE"[^>]*data-num="bn"', fdef) or re.search(r'data-num="bn"[^>]*data-f="PAGE"', fdef) or ('data-f="PAGE"' in fdef and 'data-num="bn"' in fdef), 'the page number comes back as a field, in Bangla digits', fdef)
        ok('Even pages: Office order' in hf.get('header', {}).get('even', ''), 'the even pages\' header comes back')
        await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(700)
        sh = await st(SHEETS)
        ok('Ministry of Planning' in sh[0][0] and sh[1][0] == 'H:Even pages: Office order 12/2026' and sh[1][1] == 'F:পৃষ্ঠা ২' and sh[2][0] == 'H:Office order 12/2026', 'opened, its pages show the same headers and footers', sh[:3])
        # LibreOffice reads our file the same way
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, 'hf.docx'); open(src, 'wb').write(base64.b64decode(b64))
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', td, src], capture_output=True, timeout=120)
            pdf = os.path.join(td, 'hf.pdf')
            if os.path.exists(pdf):
                txt = pdf_pages_text(pdf)
                ok('Ministry of Planning' in txt[0] and 'Office order' not in txt[0], 'LibreOffice: page 1 has the letterhead only', txt[0][:300])
                ok('Even pages: Office order 12/2026' in txt[1] and 'Office order 12/2026' in txt[2] and 'Even pages' not in txt[2], 'LibreOffice: even and odd pages have their headers', [txt[1][:120], txt[2][:120]])
                ok(re.search(r'পৃষ্ঠা\s+\S', txt[1]) is not None, 'LibreOffice: the footer has the page number', txt[1][-200:])
            else: ok(False, 'LibreOffice converted the file')

        print('7. A Word-made letter')
        await st(OPEN_DOCX, word_made())
        await page.wait_for_timeout(700)
        r = await st("() => ({ hf: __Q.S.doc.hf, first: __Q.S.doc.firstDiff, even: __Q.S.doc.evenDiff })")
        ok(r['first'] and r['even'], 'Different First Page and Odd & Even read from the file', [r['first'], r['even']])
        ok('Bangladesh Bank' in r['hf']['header']['first'] and '<table' in r['hf']['header']['first'], 'the first page\'s letterhead table', r['hf']['header']['first'][:200])
        ok('data-f="PAGE"' in r['hf']['footer']['def'] and 'Page ' in r['hf']['footer']['def'] and '>3<' not in r['hf']['footer']['def'], 'a PAGE field (begin / separate / end) becomes a page number field', r['hf']['footer']['def'])
        cells = await st("() => { const t = document.createElement('div'); t.innerHTML = __Q.S.doc.hf.header.def; return [...t.querySelectorAll('table.hf-grid td p')].map(p => p.style.textAlign + ':' + p.textContent); }")
        ok(cells == ['left:Bangladesh Bank', 'center:Circular No. 7', 'right:continued'], 'a line set out with the Header style\'s centre and right tabs becomes a three-cell table', cells)
        sh = await st(SHEETS)
        ok('Bangladesh Bank' in sh[0][0] and sh[1] == ['H:Even page header text', 'F:Even page 2'] and sh[2] == ['H:Bangladesh Bank Circular No. 7 continued', 'F:Page 3'], 'its pages: letterhead, then even and odd headers and numbered footers', sh[:3])

        print('8. A floating logo in the header')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Float', html: h }))", HTML)
        await page.wait_for_timeout(300)
        await st("""png => { const d = __Q.S.doc; d.hf.header.def = `<p style="margin:0"><img class="img-front" src="data:image/png;base64,${png}" style="width:48px;height:48px;left:-70px;top:-20px">Head office</p>`; __Q.HF.refresh(); }""", PNG)
        await page.wait_for_timeout(300)
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        hx = [z.read(n).decode() for n in z.namelist() if re.match(r'word/header\d+\.xml', n)][0]
        m = re.search(r'<wp:positionH relativeFrom="column"><wp:posOffset>(-?\d+)</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>(-?\d+)</wp:posOffset>', hx)
        ok(m and abs(int(m.group(1)) / 9525 + 70) < 1.5 and abs(int(m.group(2)) / 9525 + 20) < 1.5, 'written in front of the text, from the header\'s column and paragraph', m and m.groups())
        await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(700)
        r = await st("() => { const h = __Q.S.doc.hf.header.def; const t = document.createElement('div'); t.innerHTML = h; const im = t.querySelector('img'); return im && [im.className, parseFloat(im.style.left), parseFloat(im.style.top), im.hasAttribute('data-fx')]; }")
        ok(r and r[0] == 'img-front' and abs(r[1] + 70) < 1.5 and abs(r[2] + 20) < 1.5 and not r[3], 'opened again, it is in front of the text in the same place', r)

        print('9. The PDF and the web page')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(300)
        await st("() => __Q.HF.useDesign('bn')")
        await page.keyboard.type('পরিকল্পনা মন্ত্রণালয়')
        await st("() => __Q.HF.close()")
        await st("() => __Q.HF.pageNumbers('footer', 'right', 'of')")
        await page.wait_for_timeout(500)
        r = await st("""async () => { await document.fonts.ready; const blob = await __Q.buildVectorPdf({ title: 'Office order' }); const u = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < u.length; i += 0x8000) s += String.fromCharCode.apply(null, u.subarray(i, i + 0x8000)); return btoa(s); }""")
        open('shots/hf1.pdf', 'wb').write(base64.b64decode(r))
        txt = pdf_pages_text('shots/hf1.pdf')
        ok('পরিকল্পনা মন্ত্রণালয়' in txt[0] and '[Organization name]' in txt[0], 'the PDF: page 1 has the Bangla and English letterhead', txt[0][:300])
        ok('Page 2 of %d' % len(txt) in txt[1] and 'Page 1 of' not in txt[0], 'the PDF: page 2 is numbered "Page 2 of N"; page 1 isn\'t', txt[1][-200:])
        imgs = subprocess.run(['pdfimages', '-list', 'shots/hf1.pdf'], capture_output=True, text=True).stdout
        ok(re.search(r'^\s*1\s+\d+\s+image', imgs, re.M) is not None, 'the PDF: the logo is drawn on page 1', imgs[:400])
        html = await st("() => __Q.toStandaloneHtml()")
        ok('<header class="hfx doc hdr">' in html and 'পরিকল্পনা মন্ত্রণালয়' in html and '@top-left' not in html, 'the web page has the letterhead at the top of the page', re.findall(r'<header[^>]*>', html))

        print('10. A big logo in the account')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Big logo', html: h }))", HTML)
        await page.wait_for_timeout(300)
        r = await st("""async () => {
          const cv = document.createElement('canvas'); cv.width = 330; cv.height = 330; const x = cv.getContext('2d'); const im = x.createImageData(330, 330);
          for (let i = 0; i < im.data.length; i++) im.data[i] = (i % 4 === 3) ? 255 : Math.floor(Math.random() * 256);
          x.putImageData(im, 0, 0); const url = cv.toDataURL('image/png');
          __Q.change(() => { __Q.S.doc.hf.header.def = `<p style="margin:0"><img src="${url}" style="width:60px;height:60px"> Big logo office</p>`; }, { full: true });
          __Q.HF.refresh();
          await __Q.Saver.flush();
          const id = __Q.S.doc.id;
          const key = [...window.__mockStore.keys()].find(k => k.endsWith('/docs/' + id));
          const meta = window.__mockStore.get(key);
          const back = await __Q.CloudStore.load(id);
          return { url: url.length, cloud: __Q.Store.cloud, body: new TextEncoder().encode(JSON.stringify(meta)).length, hfChunks: (meta.hfChunks || []).length, same: back.hf.header.def === __Q.S.doc.hf.header.def, state: document.querySelector('#saveState, .save-state') && document.querySelector('#saveState, .save-state').textContent };
        }""")
        ok(r['cloud'] and r['url'] > 300000 and r['body'] < 256 * 1024 and r['hfChunks'] >= 2 and r['same'], 'a logo bigger than a stored record is kept in parts, and comes back whole', r)

        print('11. Phone width')
        b2, page2, errors2 = await new_page(p, width=400, height=860, mock_caps=True, gfonts=True)
        await load(page2)
        await page2.evaluate("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page2.wait_for_timeout(300)
        await page2.evaluate("() => __Q.HF.useDesign('centre')")
        await page2.wait_for_timeout(400)
        r = await page2.evaluate("() => ({ sw: document.documentElement.scrollWidth, cw: document.documentElement.clientWidth, tab: __Q.S.tab, ed: !!__Q.HF.editing })")
        ok(r['sw'] <= r['cw'] and r['tab'] == 'hf' and r['ed'], 'on a phone the header opens with its tab and the page doesn\'t scroll sideways', r)
        await page2.screenshot(path='shots/hf1-phone.png')
        for e in errors2: print('ERR(phone)', e)
        await b2.close()

        await page.screenshot(path='shots/hf1.png')
        errs = [e for e in errors if 'favicon' not in e[1]]
        ok(not errs, 'no errors in the console', errs[:5])
        await b.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')

asyncio.run(main())
