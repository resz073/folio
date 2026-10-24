"""Bijoy. The Bijoy keyboard typing Unicode (Ctrl+Alt+B, F12, the status bar's keyboard menu,
Review › Bijoy with its typing order, Options): ি ে ৈ before the letter and the reph after it as in
Bijoy, or reading order (UniJoy); g joins letters and makes vowels; Backspace, undo, Track Changes,
text boxes outside the page; the layout dialog. Bijoy's SutonnyMJ text to Unicode and back: the
converters on real SutonnyMJ text and a word list, Review › Convert on a selection, on table cells
and on the whole document (undo), pasting (Word's HTML by class or style, plain text by its codes,
Text Only, Keep Bijoy, the option off), opening Word files (python-docx: Normal, a heading, a
table, a header and runs split inside a syllable), web pages and text files, Unicode to Bijoy
saved to Word in SutonnyMJ; read-only; phone width."""
import asyncio, sys, json, base64, io, zipfile, re, os
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)[:900]))
    if not cond: fails.append(what)

def norm(s):
    return (s.replace('ড়', 'ড়').replace('ঢ়', 'ঢ়').replace('য়', 'য়')
             .replace('ো', 'ো').replace('ৌ', 'ৌ'))

# real SutonnyMJ text and what it says (from workbooks, news pages and letters)
GOLDEN = [
    ['Avwg evsjvq Mvb MvB', 'আমি বাংলায় গান গাই'],
    ['c~e© AvMvbMi, †KivYxMÄ, XvKv|', 'পূর্ব আগানগর, কেরাণীগঞ্জ, ঢাকা।'],
    ['cÖwZôv‡bi bvg', 'প্রতিষ্ঠানের নাম'],
    ['wmwÏK Mv‡g©›Um', 'সিদ্দিক গার্মেন্টস'],
    ['A‡±vei -2025', 'অক্টোবর -২০২৫'],
    ['¯^Y©v Mv‡g©›Um', 'স্বর্ণা গার্মেন্টস'],
    ['wewìs I Zjv b¤^i', 'বিল্ডিং ও তলা নম্বর'],
    ['wemwgjøvwni ivn&gvwbi ivwng', 'বিসমিল্লাহির রাহ্‌মানির রাহিম'],
    ['bRiæj', 'নজরুল'],
    ['‡deªæqvix -2026', 'ফেব্রুয়ারী -২০২৬'],
    ['d‡qR Avn‡¤§`', 'ফয়েজ আহম্মেদ'],
    ['MYcÖRvZš¿x evsjv‡`k miKvi', 'গণপ্রজাতন্ত্রী বাংলাদেশ সরকার'],
    ['Dc‡Rjv wbe©vnx Awdmvi', 'উপজেলা নির্বাহী অফিসার'],
    ['¯§viK bs', 'স্মারক নং'],
    ['m‡e©v”P', 'সর্বোচ্চ'],
    ['mswkøóiv Dcw¯’Z _vKvi K_v wQj|', 'সংশ্লিষ্টরা উপস্থিত থাকার কথা ছিল।'],
    ['mvew¯Œckb-wbf©i e¨emvwqK g‡W‡ji', 'সাবস্ক্রিপশন-নির্ভর ব্যবসায়িক মডেলের'],
    ['Kywgjøv AÂ‡ji ïé-K‡ii 60 kZvs‡ki', 'কুমিল্লা অঞ্চলের শুল্ক-করের ৬০ শতাংশের'],
    ['‡`‡ki iæMœwk‡íi Dbœqb', 'দেশের রুগ্নশিল্পের উন্নয়ন'],
    ['i¨ve', 'র‍্যাব'],
    ['‡KŠkj ˆZwi', 'কৌশল তৈরি'],
]
WORDS = ['আমি', 'বাংলায়', 'গান', 'গাই', 'কর্মকর্তা', 'স্বাক্ষর', 'প্রজাতন্ত্র', 'মন্ত্রণালয়', 'উপজেলা', 'নির্বাহী', 'জেলা', 'প্রশাসক', 'ক্ষমতা',
         'বিজ্ঞপ্তি', 'অফিস', 'আদেশ', 'স্মারক', 'তারিখ', 'সংযুক্তি', 'অনুলিপি', 'সদয়', 'জ্ঞাতার্থে', 'কার্যার্থে', 'ব্যবস্থা', 'গ্রহণের', 'জন্য',
         'অনুরোধ', 'করা', 'হলো', 'রাষ্ট্রপতির', 'আদেশক্রমে', 'সচিব', 'বাংলাদেশ', 'শিক্ষা', 'স্বাস্থ্য', 'পরিবার', 'কল্যাণ', 'শৃঙ্খলা', 'উজ্জ্বল',
         'সৌজন্য', 'ঐতিহ্য', 'ঔষধ', 'ঋণ', 'পৌঁছানো', 'চাঁদপুর', 'র‍্যাব', 'অ্যাকাউন্ট', 'সংক্ষিপ্ত', 'দ্বন্দ্ব', 'উদ্ধৃতি', 'ইঙ্গিত', 'পঙ্‌ক্তি']

def word_fixture():
    """a letter as Bijoy users make it: Normal in SutonnyMJ, an English run in Times New Roman, a
    heading in SutonnyMJ, a table, a header, and a word whose ে is bold (its own run)"""
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
    d = docx.Document()
    def mj(style):
        style.font.name = 'SutonnyMJ'
        rf = style.element.rPr.rFonts
        rf.set(qn('w:hAnsi'), 'SutonnyMJ')
    mj(d.styles['Normal'])
    mj(d.styles['Heading 1'])
    d.add_heading('MYcÖRvZš¿x evsjv‡`k miKvi', level=1)
    p = d.add_paragraph('¯§viK bs: 05.12 ')
    r = p.add_run('Office of the DC'); r.font.name = 'Times New Roman'
    p2 = d.add_paragraph()
    r = p2.add_run('‡'); r.bold = True
    p2.add_run('Kv‡bv Kg©KZ©v Dcw¯’Z _vK‡eb|')
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = 'µwgK'
    t.cell(0, 1).text = 'bvg'
    d.sections[0].header.paragraphs[0].text = 'Rvbyqvwi 2026'
    bio = io.BytesIO(); d.save(bio)
    return base64.b64encode(bio.getvalue()).decode()

OPEN_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); r.title = 'Imported'; await __Q.openImported(r); return { bijoy: r.settings.bijoy || 0, font: r.settings.font, html: r.html, hdr: r.settings.hf && r.settings.hf.header.def }; }"""
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""
TXT = "() => document.getElementById('editor').innerText.replace(/\\n+/g, '\\n').trim()"

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1366, height=900, mock_caps=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)
        async def fresh(html='<p><br></p>'):
            await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Bijoy', html: h }))", html)
            await page.wait_for_timeout(250)
            await page.click('#editor')
            await st("() => { const e = document.getElementById('editor'); const r = document.createRange(); r.selectNodeContents(e.lastElementChild); r.collapse(false); __Q.setRange(r); }")

        print('1. The converters')
        res = await st("g => g.map(([a, u]) => __Q.bijoyToUnicode(a))", GOLDEN)
        bad = [[g[0], r, g[1]] for g, r in zip(GOLDEN, res) if norm(r) != norm(g[1])]
        ok(not bad, f'real SutonnyMJ text reads right ({len(GOLDEN)} lines: conjuncts, reph, ি ে ৈ before their letters, ো ৌ, ঁ, র‍্য)', bad)
        rt = await st("w => w.map(x => [x, __Q.bijoyFromUnicode(x), __Q.bijoyToUnicode(__Q.bijoyFromUnicode(x))])", WORDS)
        bad = [r for r in rt if norm(r[2]) != norm(r[0]) and norm(r[2]).replace('্‌', '্') != norm(r[0]).replace('্‌', '্')]
        ok(not bad, f'Unicode → Bijoy → Unicode gives the words back ({len(WORDS)} words)', bad)
        sample = {r[0]: r[1] for r in rt}
        ok(sample['আমি'] == 'Avwg' and sample['কর্মকর্তা'] == 'Kg©KZ©v' and sample['ক্ষমতা'] == '¶gZv' and sample['স্মারক'] == '¯§viK' and sample['জেলা'] == '‡Rjv' and sample['রাষ্ট্রপতির'] == 'ivóªcwZi',
           'Unicode → Bijoy writes what Bijoy types (Avwg, Kg©KZ©v, ¶gZv, ¯§viK, ‡Rjv, ivóªcwZi)', [sample['আমি'], sample['কর্মকর্তা'], sample['ক্ষমতা'], sample['স্মারক'], sample['জেলা'], sample['রাষ্ট্রপতির']])
        mixed = await st("() => __Q.bijoyToUnicode('শিল্প মন্ত্রী eZ©gv‡b gqgbwmsn')")
        ok(mixed == 'শিল্প মন্ত্রী বর্তমানে ময়মনসিংহ', 'Bangla already in Unicode stays as it is beside Bijoy text', mixed)
        det = await st("() => [['Avwg evsjvq Mvb MvB', 'Dc‡Rjv wbe©vnx Awdmvi', 'cÖwZôv‡bi bvg', 'Avãyj gÄy'].map(s => __Q.bijoyScore(s).bijoy), ['The office will remain closed on Sunday.', 'Grand Total', 'MacBook iPhone', 'café au lait', 'Memo No: 05.12.0000.000.11'].map(s => __Q.bijoyScore(s).bijoy)]")
        ok(det == [[True] * 4, [False] * 5], 'Bijoy text is told from English by its codes', det)
        fonts = await st("() => ['SutonnyMJ', 'Sutonny MJ', 'SulekhaMJ', 'NikoshBAN', 'Kalpurush ANSI', 'SutonnyOMJ', 'Nikosh', 'Times New Roman'].map(__Q.isBijoyFont)")
        ok(fonts == [True] * 5 + [False] * 3, 'Bijoy fonts (SutonnyMJ, SulekhaMJ, NikoshBAN, Kalpurush ANSI), not SutonnyOMJ or Nikosh', fonts)

        print('2. Typing with the Bijoy keyboard')
        await fresh()
        await page.keyboard.press('Control+Alt+b')
        await page.wait_for_timeout(100)
        s = await st("() => [__Q.S.bijoy, __Q.S.avro, document.getElementById('sbAvro').textContent, __Q.Prefs.bnKbd, __Q.Prefs.bijoy]")
        ok(s == [True, False, 'ক Bijoy', 'bijoy', True], 'Ctrl+Alt+B turns the Bijoy keyboard on; the status bar says ক Bijoy', s)
        await page.keyboard.type('gfdm hfQVfW ofb ofgdG')
        t = await st(TXT)
        ok(norm(t) == norm('আমি বাংলায় গান গাই।'), 'g f = আ, d m = মি (ি typed first), G = ।: আমি বাংলায় গান গাই।', t)
        await page.keyboard.press('Enter')
        await page.keyboard.type('jmA cjf djgNA hgV Ff vZf')
        t2 = await st("() => document.getElementById('editor').lastElementChild.textContent")
        ok(norm(t2) == norm('কর্ম কো র্ক্ষি ব্ল আ র‍্যা'), 'reph after its letter (j m A = কর্ম), c j f = কো, d j g N A = র্ক্ষি, g joins (h g V = ব্ল), F f = আ, v Z = র‍্য', t2)
        await page.keyboard.type(' jd')
        await page.keyboard.press('Backspace')
        t3 = await st("() => document.getElementById('editor').lastElementChild.textContent")
        ok(t3.endswith(' ক'), 'Backspace takes back the last key (j d → ক)', t3)
        await page.keyboard.type('m')
        t4 = await st("() => document.getElementById('editor').lastElementChild.textContent")
        ok(t4.endswith(' কম'), 'and typing goes on from there', t4)
        await page.keyboard.press('F12')
        await page.keyboard.type(' ok')
        s = await st("() => [__Q.S.bijoy, document.getElementById('sbAvro').textContent, document.getElementById('editor').lastElementChild.textContent.slice(-3)]")
        ok(s == [False, 'ENG', ' ok'], 'F12 turns it off: English again', s)
        await page.keyboard.press('F12')
        ok(await st("() => [__Q.S.bijoy, __Q.S.avro]") == [True, False], 'F12 turns on the Bangla keyboard last chosen (Bijoy)')
        await page.keyboard.press('Control+z')
        await page.wait_for_timeout(100)
        ok('ok' not in await st(TXT), 'Undo takes typing back')
        # reading order (UniJoy)
        await st("() => __Q.Bijoy.setOrder('unicode')")
        await fresh()
        await page.keyboard.type('jd jAm hfQVf')
        t = await st(TXT)
        ok(norm(t) == norm('কি কর্ম বাংলা'), 'in reading order (UniJoy): j d = কি, j A m = কর্ম', t)
        await st("() => __Q.Bijoy.setOrder('bijoy')")
        ok(await st("() => __Q.Prefs.bijoyOrder") == 'bijoy', 'the order is kept in the settings')
        # Track Changes
        await fresh('<p>Note: </p>')
        await st("() => __Q.Track.toggle()")
        await page.keyboard.type('jfwA')
        tc = await st("() => [...document.querySelectorAll('#editor ins.tc')].map(e => e.textContent).join('')")
        ok(norm(tc) == norm('কার্য'), 'with Track Changes on, the Bangla typed is a tracked insertion (j f w A = কার্য)', tc)
        await st("() => __Q.Track.toggle()")
        # a text box outside the page: Find
        await page.keyboard.press('Control+f')
        await page.wait_for_timeout(200)
        await page.keyboard.type('ofb')
        fq = await st("() => document.getElementById('findQ').value")
        ok(norm(fq) == norm('গান'), 'boxes outside the page (Find) type with Bijoy too', fq)
        await page.keyboard.press('Escape')
        # the status bar's keyboard menu
        await page.click('#sbKbd')
        await page.wait_for_timeout(150)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(e => [e.querySelector('.lb').textContent, e.getAttribute('aria-checked')])")
        ok(items[:3] == [['English', 'false'], ['Avro Phonetic', 'false'], ['Bijoy', 'true']] and any('UniJoy' in i[0] for i in items), 'the status bar\'s arrow: English, Avro Phonetic, Bijoy (checked), the typing order, the layout', items)
        await page.locator('.pop .mi', has_text='Avro Phonetic').click()
        await page.wait_for_timeout(150)
        s = await st("() => [__Q.S.avro, __Q.S.bijoy, document.getElementById('sbAvro').textContent, __Q.Prefs.bnKbd]")
        ok(s == [True, False, 'অ Avro', 'avro'], 'choosing Avro Phonetic there turns Bijoy off', s)
        await page.keyboard.press('Control+Alt+b')
        ok(await st("() => [__Q.S.avro, __Q.S.bijoy]") == [False, True], 'Ctrl+Alt+B goes from Avro to Bijoy')
        # Review › Bijoy
        await page.click('#tab-review')
        await page.wait_for_timeout(150)
        pressed = await st("() => { const b = [...document.querySelectorAll('.rb')].find(x => x.offsetParent && x.querySelector('.cap') && x.querySelector('.cap').textContent === 'Bijoy'); return b && b.getAttribute('aria-pressed'); }")
        ok(pressed == 'true', 'Review › Bangla › Bijoy shows it is on', pressed)
        await page.locator('.rsplit:visible:has(.cap:text-is("Bijoy")) .rb.arrow').click()
        await page.wait_for_timeout(150)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(e => e.querySelector('.lb').textContent)")
        ok(items == ['ি ে ৈ before the letter (Bijoy)', 'Reading order (UniJoy)', 'Bijoy Keyboard Layout…'], 'its arrow: the typing order and the layout', items)
        await page.locator('.pop .mi', has_text='Bijoy Keyboard Layout').click()
        await page.wait_for_timeout(300)
        lay = await st("() => { const d = document.querySelector('.dialog'); return d && [d.querySelector('h2').textContent, d.querySelectorAll('.bj-key').length, d.querySelector('.bj-key.link .lo').textContent, d.querySelector('#bjOrder').value, [...d.querySelectorAll('.bj-key')].find(k => k.querySelector('.k').textContent === 'J').querySelector('.lo').textContent]; }")
        ok(lay == ['Bijoy Keyboard', 47, '◌্', 'bijoy', 'ক'], 'the layout: 47 keys with their letters (J is ক), g the joining key, the typing order', lay)
        await page.screenshot(path='shots/bj1-layout.png')
        await page.keyboard.press('Escape')
        # Options › Bangla
        await st("() => __Q.Backstage.open('options')")
        await page.wait_for_timeout(250)
        o = await st("() => [document.getElementById('opt-avro').checked, document.getElementById('opt-bijoy').checked, document.getElementById('optF12').value, document.getElementById('optBjOrder').value, document.getElementById('opt-bijoy-auto').checked]")
        ok(o == [False, True, 'bijoy', 'bijoy', True], 'Options › Bangla: Avro, Bijoy, what F12 turns on, the typing order, converting Bijoy text', o)
        await page.click('label:has(#opt-avro)')
        await page.wait_for_timeout(100)
        o = await st("() => [document.getElementById('opt-avro').checked, document.getElementById('opt-bijoy').checked, __Q.S.avro, __Q.S.bijoy]")
        ok(o == [True, False, True, False], 'turning Avro on there turns Bijoy off (and its box)', o)
        await st("() => __Q.Backstage.close()")
        await st("() => __Q.Bijoy.use('en')")

        print('3. Pasting Bijoy text')
        await fresh()
        word_html = ('<html><head><style><!-- p.MsoNormal, li.MsoNormal, div.MsoNormal {font-size:14.0pt; font-family:SutonnyMJ;} --></style></head><body>'
                     '<!--StartFragment--><p class=MsoNormal>Avwg evsjvq Mvb MvB <span style="font-family:\'Times New Roman\'">(Folio)</span> e‡l©</p>'
                     '<p><span style="font-family:SutonnyMJ;mso-bidi-font-family:SutonnyMJ">wb‡`©kbv †gvZv‡eK</span> ok</p><!--EndFragment--></body></html>')
        await st("([h]) => { __Q.Paste.last = { html: h, text: 'x' }; __Q.doPaste('keep', h, 'x', true); }", [word_html])
        await page.wait_for_timeout(200)
        pz = await st("() => [document.getElementById('editor').innerText.replace(/\\n+/g, '\\n').trim(), document.getElementById('editor').innerHTML.includes('SutonnyMJ'), [...document.querySelectorAll('#toastRoot .toast')].map(t => t.textContent).pop()]")
        ok(norm(pz[0]) == norm('আমি বাংলায় গান গাই (Folio) বর্ষে\nনির্দেশনা মোতাবেক ok') and not pz[1], 'Word\'s HTML in SutonnyMJ (by class and by style) comes in as Unicode; English in Times New Roman stays', pz[:2])
        ok(pz[2] and 'converted to Unicode' in pz[2] and 'Keep Bijoy' in pz[2], 'a note says so, with Keep Bijoy', pz[2])
        await page.locator('#toastRoot .toast button', has_text='Keep Bijoy').last.click()
        await page.wait_for_timeout(250)
        kb = await st("() => [document.getElementById('editor').innerText.replace(/\\n+/g, '\\n').trim(), document.getElementById('editor').innerHTML.includes('SutonnyMJ')]")
        ok(kb[0].startswith('Avwg evsjvq') and kb[1], 'Keep Bijoy pastes it again as it was, in SutonnyMJ', kb)
        await st("() => __Q.PasteChip.redo('keep', { noBijoy: false })")
        await page.wait_for_timeout(200)
        ok(norm(await st(TXT)).startswith(norm('আমি বাংলায়')), 'and it can be converted again')
        # plain text
        await fresh()
        await st("() => { __Q.Paste.last = null; __Q.doPaste('keep', '', 'Dc‡Rjv wbe©vnx Awdmvi Ges mswkøó mKj‡K Rvbv‡bv hv‡”Q', false); }")
        t = await st(TXT)
        ok(norm(t) == norm('উপজেলা নির্বাহী অফিসার এবং সংশ্লিষ্ট সকলকে জানানো যাচ্ছে'), 'plain text in Bijoy codes is known by them and converted', t)
        await fresh()
        await st("() => __Q.doPaste('keep', '', 'Please find the attached report for the quarter.', false)")
        ok(await st(TXT) == 'Please find the attached report for the quarter.', 'English plain text stays English')
        await fresh()
        await st("([h]) => __Q.doPaste('text', h, 'Avwg evsjvq Mvb MvB', false)", [word_html])
        t = await st(TXT)
        ok(norm(t).startswith(norm('আমি বাংলায় গান গাই (Folio)')), 'Paste › Text Only of Word\'s HTML: the SutonnyMJ text still converts (English stays)', t)
        await st("() => __Q.Bijoy.setAuto(false)")
        await fresh()
        await st("([h]) => __Q.doPaste('keep', h, 'x', false)", [word_html])
        ok((await st(TXT)).startswith('Avwg evsjvq'), 'with the option off, pasted Bijoy text stays as it is')
        await st("() => __Q.Bijoy.setAuto(true)")

        print('4. Review › Convert')
        await fresh('<p id="a"><span style="font-family: SutonnyMJ">Avwg evsjvq Mvb MvB</span> and English</p><p id="b">উপজেলা নির্বাহী অফিসার, Dhaka</p><p id="c">Plain English line.</p>')
        await st("() => { const p = document.getElementById('a'); const r = document.createRange(); r.selectNodeContents(p); __Q.setRange(r); }")
        await page.click('#tab-review')
        await page.locator('.rb:visible:has(.cap:text-is("Convert"))').click()
        await page.wait_for_timeout(150)
        items = await st("() => [...document.querySelectorAll('.pop .mi')].map(e => [e.querySelector('.lb').textContent, e.getAttribute('aria-checked')])")
        ok(items == [['Bijoy to Unicode', None], ['Unicode to Bijoy (SutonnyMJ)', None], ['Convert Bijoy text when pasting and opening Word files', 'true']], 'Review › Convert: Bijoy to Unicode, Unicode to Bijoy, the paste/open option', items)
        await page.locator('.pop .mi', has_text='Bijoy to Unicode').click()
        await page.wait_for_timeout(200)
        a = await st("() => { const p = document.getElementById('a'); return [p.textContent, p.innerHTML.includes('SutonnyMJ'), window.getSelection().toString()]; }")
        ok(norm(a[0]) == norm('আমি বাংলায় গান গাই and English') and not a[1], 'Bijoy to Unicode on a selection: its SutonnyMJ text converts and loses the font; the English beside it stays', a)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(150)
        ok(await st("() => document.getElementById('a').textContent") == 'Avwg evsjvq Mvb MvB and English', 'Undo brings the Bijoy text back')
        await st("() => { const p = document.getElementById('b'); const r = document.createRange(); r.selectNodeContents(p); __Q.setRange(r); }")
        await st("() => __Q.Bijoy.convert('bijoy')")
        await page.wait_for_timeout(150)
        bj = await st("() => { const p = document.getElementById('b'); const sp = p.querySelector('span'); return [p.textContent, sp && getComputedStyle(sp).fontFamily.split(',')[0].replace(/\"/g, '')]; }")
        ok(bj == ['Dc‡Rjv wbe©vnx Awdmvi, Dhaka', 'SutonnyMJ'], 'Unicode to Bijoy: the Bangla becomes Bijoy text in SutonnyMJ; English stays', bj)
        await st("() => { const p = document.getElementById('b'); const r = document.createRange(); r.selectNodeContents(p); __Q.setRange(r); }")
        await st("() => __Q.Bijoy.convert('unicode')")
        await page.wait_for_timeout(150)
        back = await st("() => [document.getElementById('b').textContent, document.getElementById('b').innerHTML.includes('SutonnyMJ')]")
        ok(norm(back[0]) == norm('উপজেলা নির্বাহী অফিসার, Dhaka') and not back[1], 'and Bijoy to Unicode brings it back the same (only the SutonnyMJ part: Dhaka stays)', back)
        # whole document, the document's font in SutonnyMJ
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Old letter', font: 'SutonnyMJ', html: h }))", '<p id="x">Avwg evsjvq Mvb MvB</p><p id="y"><span style="font-family: Arial">Reference: 12/3</span></p><p>wb‡`©kbv</p>')
        await page.wait_for_timeout(250)
        await page.click('#editor')
        await st("() => { const r = document.createRange(); r.setStart(document.getElementById('x').firstChild, 2); r.collapse(true); __Q.setRange(r); }")
        await st("() => __Q.Bijoy.convert('unicode')")
        await page.wait_for_timeout(200)
        w = await st("() => [document.getElementById('editor').innerText.replace(/\\n+/g, '\\n').trim(), __Q.S.doc.font]")
        ok(norm(w[0]) == norm('আমি বাংলায় গান গাই\nReference: 12/3\nনির্দেশনা') and w[1] == 'Times New Roman', 'with nothing selected, all the text in a Bijoy font converts (the document\'s own font too: English then goes to Times New Roman); Arial text stays', w)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(150)
        ok(await st("() => [document.getElementById('x').textContent, __Q.S.doc.font]") == ['Avwg evsjvq Mvb MvB', 'SutonnyMJ'], 'one Undo puts all of it back, the font too')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'English', html: h }))", '<p>Just English.</p>')
        await page.wait_for_timeout(200)
        await page.click('#editor')
        await st("() => __Q.Bijoy.convert('unicode')")
        await page.wait_for_timeout(100)
        tz = await st("() => [...document.querySelectorAll('#toastRoot .toast')].map(t => t.textContent).pop()")
        ok(tz and 'Nothing here is in a Bijoy font' in tz and await st(TXT) == 'Just English.', 'nothing in a Bijoy font: a note, and nothing changes', tz)
        # table cells
        await fresh('<table class="tbl-grid"><tbody><tr><td><p>µwgK</p></td><td><p>bvg</p></td></tr><tr><td><p>1</p></td><td><p>Kwig</p></td></tr></tbody></table><p>after</p>')
        await st("() => { const t = document.querySelector('#editor table'); __Q.Tables.setRect(t.rows[0].cells[0], t.rows[1].cells[1]); }")
        ok(await st("() => !!__Q.Tables.active()"), 'four cells selected')
        await st("() => __Q.Bijoy.convert('unicode')")
        await page.wait_for_timeout(150)
        cells = await st("() => [...document.querySelectorAll('#editor td')].map(c => c.textContent)")
        ok([norm(c) for c in cells] == [norm('ক্রমিক'), norm('নাম'), '১', norm('করিম')], 'selected table cells convert, each one', cells)

        print('5. Word files')
        b64 = word_fixture()
        r = await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(400)
        body = await st(TXT)
        ok(r['bijoy'] >= 5 and r['font'] == 'Times New Roman', 'a Word file in SutonnyMJ: its Bijoy text is converted as it opens (Normal\'s font becomes Times New Roman for the English)', [r['bijoy'], r['font']])
        ok('SutonnyMJ' not in r['html'] and 'data-bj' not in r['html'], 'nothing is left in SutonnyMJ, and no marks of the reading', 'SutonnyMJ' in r['html'])
        lines = [norm(x) for x in body.split('\n') if x.strip()]
        want = [norm(x) for x in ['গণপ্রজাতন্ত্রী বাংলাদেশ সরকার', 'স্মারক নং: ০৫.১২ Office of the DC', 'কোনো কর্মকর্তা উপস্থিত থাকবেন।']]
        ok(lines[:3] == want, 'the heading, the memo line (with its English run kept) and a word whose ে is bold on its own read right', lines[:5])
        cells = await st("() => [...document.querySelectorAll('#editor td')].map(c => c.textContent)")
        ok([norm(c) for c in cells] == [norm('ক্রমিক'), norm('নাম')], 'the table\'s cells', cells)
        hd = await st("() => [...document.querySelectorAll('#sheets .sheet .hfx.hdr')].map(x => x.textContent)[0]")
        ok(hd and norm(hd) == norm('জানুয়ারি ২০২৬'), 'the header too', hd)
        empty = await st("() => [...document.querySelectorAll('#editor span, #editor b, #editor strong')].filter(x => !x.textContent && !x.querySelector('img, br')).length")
        ok(empty == 0, 'the bold ে went to its letter (a sign can\'t be bold apart from it): no empty bold run is left', empty)
        await st("() => __Q.Bijoy.setAuto(false)")
        r2 = await st(OPEN_DOCX, b64)
        ok(r2['bijoy'] == 0 and 'SutonnyMJ' in r2['html'] + r2['font'], 'with the option off it opens as it is, in SutonnyMJ', [r2['bijoy'], r2['font']])
        await st("() => __Q.Bijoy.setAuto(true)")
        # Unicode → Bijoy saved to Word
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'To Bijoy', html: h }))", '<p>সভার কার্যবিবরণী</p><p>Meeting minutes</p>')
        await page.wait_for_timeout(200)
        await page.click('#editor')
        await st("() => __Q.Bijoy.convert('bijoy')")
        await page.wait_for_timeout(150)
        d64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(d64)))
        x = z.read('word/document.xml').decode('utf8')
        ok('w:ascii="SutonnyMJ"' in x and 'mfvi Kvh©weeiYx' in x and 'Meeting minutes' in x, 'converted to Bijoy, it goes to Word in SutonnyMJ, as Bijoy users have it', re.findall(r'<w:t[^>]*>[^<]*</w:t>', x)[:4])
        r3 = await st(OPEN_DOCX, d64)
        t = await st(TXT)
        ok(norm(t.split('\n')[0]) == norm('সভার কার্যবিবরণী'), 'and that file opens back into Unicode', t)

        print('6. Web pages and text files')
        h = '<html><body><p style="font-family: SutonnyMJ">Avgv‡`i Awdm</p><p>English line</p></body></html>'
        r = await st("""async (h) => { const f = new File([h], 'old.html', { type: 'text/html' }); const r = await __Q.Files.importFile(f); return [r.html, r.settings]; }""", h)
        ok('আমাদের অফিস' in norm(r[0]).replace('য়', 'য়') or norm('আমাদের অফিস') in norm(r[0]), 'a web page in SutonnyMJ opens as Unicode', r)
        r = await st("""async () => { const f = new File(['Avgv‡`i Awdm Lyj‡e iweevi\\nmKj‡K Rvbv‡bv hv‡”Q'], 'old.txt', { type: 'text/plain' }); const r = await __Q.Files.importFile(f); return [r.html, r.settings]; }""")
        ok(norm('আমাদের অফিস খুলবে রবিবার') in norm(r[0]) and r[1].get('bijoy') == 1, 'a text file in Bijoy codes too', r)
        r = await st("""async () => { const f = new File(['The office opens on Sunday.'], 'en.txt', { type: 'text/plain' }); const r = await __Q.Files.importFile(f); return [r.html, r.settings]; }""")
        ok('The office opens on Sunday.' in r[0] and not r[1].get('bijoy'), 'an English text file stays English', r)

        print('7. Read-only and phone width')
        await fresh('<p><span style="font-family: SutonnyMJ">Avwg</span></p>')
        await st("() => { __Q.S.readOnly = true; }")
        await st("() => __Q.Bijoy.convert('unicode')")
        ok(await st(TXT) == 'Avwg', 'read-only: Convert changes nothing')
        await st("() => { __Q.S.readOnly = false; }")
        errs = [e for e in errors if 'favicon' not in e[1]]
        await b.close()
        b, page, errors2 = await new_page(p, width=390, height=844, mock_caps=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)
        await page.wait_for_timeout(300)
        sb = await st("() => { const a = document.getElementById('sbAvro').getBoundingClientRect(), k = document.getElementById('sbKbd').getBoundingClientRect(); return [a.width > 20, k.width > 8, k.right <= innerWidth]; }")
        ok(sb == [True, True, True], 'phone: the keyboard button and its arrow are in the status bar', sb)
        await page.click('#sbKbd')
        await page.wait_for_timeout(150)
        m = await st("() => { const r = document.querySelector('.pop').getBoundingClientRect(); return [r.left >= 0, r.right <= innerWidth, r.top >= 0, r.bottom <= innerHeight]; }")
        ok(m == [True] * 4, 'its menu fits the screen', m)
        await page.keyboard.press('Escape')
        await st("() => __Q.Bijoy.showLayout()")
        await page.wait_for_timeout(300)
        lw = await st("() => { const d = document.querySelector('.dialog').getBoundingClientRect(), k = document.querySelector('.bj-kb'); return [d.right <= innerWidth + 1, document.documentElement.scrollWidth <= innerWidth, k.scrollWidth >= k.clientWidth]; }")
        ok(lw[:2] == [True, True], 'the layout dialog fits a phone (the keyboard scrolls inside it)', lw)
        await page.screenshot(path='shots/bj1-phone.png')
        for e in errs + [e for e in errors2 if 'favicon' not in e[1]]: print('ERR', e)
        await b.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')
asyncio.run(main())
