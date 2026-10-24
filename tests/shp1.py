"""Shapes and text boxes. Insert › Shapes (drag; Shift; a click for Word's 1-inch shape; Esc),
Insert › Text Box; typing in a shape (a letter, Enter, double-click, a text box's own click) and
one undo for its text; the handles (corners with Shift, sides of a turned shape, the round
handle, a line's ends), arrow keys; Shape Format (styles, fill, outline weight/dashes/arrows,
text alignment, fitting, size, rotate and flip, wrapping, alignment, the Format Shape dialog);
delete, copy and paste; saved and opened again; the PDF (vector and the canvas painter); Word
files out and back, LibreOffice, and shapes made by Word (2016 text box, 365 style references to
the theme, a group, an older VML text box, one in the line, one in a header); the web page,
Markdown and text; read-only; Backspace and Delete beside a shape (the paragraphs join, the shape
stays); phone width."""
import asyncio, sys, json, base64, io, zipfile, re, subprocess, os, tempfile, math
sys.path.insert(0, 'tests')
from harness import *

fails = []
def ok(cond, what, got=None):
    print(('  ok   ' if cond else '  FAIL ') + what + ('' if cond or got is None else '  -> ' + json.dumps(got, ensure_ascii=False)[:900]))
    if not cond: fails.append(what)

HTML = ('<h1>Office order</h1><p>The office will remain closed on Sunday for the national holiday.</p>'
        '<p>All staff are requested to note this and plan their work accordingly.</p>'
        '<p>Duty officers will be named in a separate order.</p><p>Copy to all sections.</p><p>End of order.</p>')
SHP = """() => [...document.querySelectorAll('#editor span.shp')].map(s => ({ cls: s.className, style: s.getAttribute('style') || '', a: Object.fromEntries([...s.attributes].filter(x => x.name.startsWith('data-')).map(x => [x.name.slice(5), x.value])), text: (s.querySelector(':scope > template') || { innerHTML: '' }).innerHTML, drawn: !!(s.shadowRoot && s.shadowRoot.querySelector('svg')) }))"""
GET_DOCX = """async () => { const blob = await __Q.buildDocx(); const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000)); return btoa(s); }"""
OPEN_DOCX = """async (b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); const r = await __Q.readDocxFile(u.buffer); r.title = 'Imported'; await __Q.openImported(r); return true; }"""
GET_PDF = """async () => { await document.fonts.ready; const blob = await __Q.buildVectorPdf({ title: 'Shapes' }); const u = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < u.length; i += 0x8000) s += String.fromCharCode.apply(null, u.subarray(i, i + 0x8000)); return btoa(s); }"""
CARET_END = """i => { const p = document.querySelectorAll('#editor p')[i]; const r = document.createRange(); r.selectNodeContents(p); r.collapse(false); getSelection().removeAllRanges(); getSelection().addRange(r); document.getElementById('editor').focus(); }"""
# a shape's box on the screen: its centre, its own size (scaled), its turn
BOX = """q => { const s = document.querySelector('#editor ' + q); const r = s.getBoundingClientRect(); return { cx: (r.left + r.right) / 2, cy: (r.top + r.bottom) / 2, w: parseFloat(s.style.width) * __Q.S.zoom, h: parseFloat(s.style.height) * __Q.S.zoom, rot: +(s.dataset.rot || 0) }; }"""

def pdf_text(path):
    return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout

def page_png(pdf, res=72):
    from PIL import Image
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(['pdftoppm', '-r', str(res), '-png', '-f', '1', '-l', '1', pdf, os.path.join(td, 'p')], check=True)
        f = [x for x in os.listdir(td) if x.endswith('.png')][0]
        return Image.open(os.path.join(td, f)).convert('RGB')

NS = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w10="urn:schemas-microsoft-com:office:word"')
A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
WORD_BOX = f'''<w:r><mc:AlternateContent {NS}><mc:Choice Requires="wps"><w:drawing><wp:anchor distT="45720" distB="45720" distL="114300" distR="114300" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>3200400</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="2360930" cy="640080"/><wp:effectExtent l="0" t="0" r="22860" b="11430"/><wp:wrapSquare wrapText="bothSides"/><wp:docPr id="217" name="Text Box 2"/><wp:cNvGraphicFramePr><a:graphicFrameLocks {A}/></wp:cNvGraphicFramePr><a:graphic {A}><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"><a:spLocks noChangeArrowheads="1"/></wps:cNvSpPr><wps:spPr bwMode="auto"><a:xfrm><a:off x="0" y="0"/><a:ext cx="2360930" cy="640080"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:miter lim="800000"/><a:headEnd/><a:tailEnd/></a:ln></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r><w:t>Memo No. 45/2026</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" anchor="t" anchorCtr="0"><a:spAutoFit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape id="Text Box 2" o:spid="_x0000_s1026" type="#_x0000_t202" style="position:absolute;margin-left:252pt;margin-top:0;width:185.9pt;height:50.4pt;z-index:251659264"><v:textbox style="mso-fit-shape-to-text:t"><w:txbxContent><w:p><w:r><w:t>Memo No. 45/2026</w:t></w:r></w:p></w:txbxContent></v:textbox><w10:wrap type="square"/></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>'''
WORD_OVAL = f'''<w:r><mc:AlternateContent {NS}><mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251660288" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>1828800</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>914400</wp:posOffset></wp:positionV><wp:extent cx="1371600" cy="914400"/><wp:effectExtent l="0" t="0" r="19050" b="19050"/><wp:wrapNone/><wp:docPr id="2" name="Oval 2"/><wp:cNvGraphicFramePr/><a:graphic {A}><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm rot="1800000"><a:off x="0" y="0"/><a:ext cx="1371600" cy="914400"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom></wps:spPr><wps:style><a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef><a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef><a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></wps:style><wps:txbx><w:txbxContent><w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:t>Seal</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" anchor="ctr" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:oval id="Oval 2" style="position:absolute;margin-left:144pt;margin-top:72pt;width:108pt;height:72pt;z-index:251660288" fillcolor="#4f81bd" strokecolor="#243f60" strokeweight="2pt"/></w:pict></mc:Fallback></mc:AlternateContent></w:r>'''
WORD_GROUP = f'''<w:r><mc:AlternateContent {NS}><mc:Choice Requires="wpg"><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251661312" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>190500</wp:posOffset></wp:positionV><wp:extent cx="1828800" cy="457200"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="3" name="Group 3"/><wp:cNvGraphicFramePr/><a:graphic {A}><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"><wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="457200"/><a:chOff x="0" y="0"/><a:chExt cx="1828800" cy="457200"/></a:xfrm></wpg:grpSpPr><wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFC000"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr><wps:bodyPr/></wps:wsp><wps:wsp><wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="914400" y="228600"/><a:ext cx="914400" cy="0"/></a:xfrm><a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="19050"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp></wpg:wgp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>'''
WORD_VML = f'''<w:r><w:pict {NS}><v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe"><v:stroke joinstyle="miter"/><v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype><v:shape id="Text Box 4" o:spid="_x0000_s1030" type="#_x0000_t202" style="position:absolute;margin-left:0;margin-top:9pt;width:162pt;height:36pt;z-index:251662336;mso-position-horizontal:right;mso-position-horizontal-relative:margin;mso-position-vertical-relative:text" strokecolor="#c00000" strokeweight="1.5pt"><v:textbox style="mso-fit-shape-to-text:t" inset="5.4pt,3.6pt,5.4pt,3.6pt"><w:txbxContent><w:p><w:r><w:t>Old Word box</w:t></w:r></w:p></w:txbxContent></v:textbox><w10:wrap type="none"/></v:shape></w:pict></w:r>'''
WORD_INLINE = f'''<w:r><mc:AlternateContent {NS}><mc:Choice Requires="wps"><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="457200" cy="228600"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:docPr id="5" name="Rectangle 5"/><wp:cNvGraphicFramePr/><a:graphic {A}><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="457200" cy="228600"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="00B050"/></a:solidFill><a:ln><a:noFill/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>'''
WORD_HDR = f'''<w:r><mc:AlternateContent {NS}><mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251663360" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="1371600" cy="320040"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="6" name="Text Box 6"/><wp:cNvGraphicFramePr/><a:graphic {A}><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1371600" cy="320040"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="6350"><a:solidFill><a:srgbClr val="1F3A5F"/></a:solidFill></a:ln></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r><w:t>Ref: 12/2026</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr lIns="91440" tIns="45720" rIns="91440" bIns="45720" anchor="ctr"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>'''

def word_fixture():
    """a Word file with shapes as Word writes them (python-docx for the rest)"""
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = 'HDRMARK'
    for i in range(1, 6): d.add_paragraph(f'MARK{i}')
    d.add_paragraph('The last paragraph.')
    bio = io.BytesIO(); d.save(bio)
    zi = zipfile.ZipFile(io.BytesIO(bio.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == 'word/document.xml':
                x = data.decode()
                for i, (sh, txt) in enumerate([(WORD_BOX, 'A paragraph with a text box set to its right.'), (WORD_OVAL, 'A paragraph with an oval in front of it.'), (WORD_GROUP, 'A group: a box and an arrow.'), (WORD_VML, 'An older Word text box.'), (WORD_INLINE, None)], 1):
                    para = f'<w:p>{sh}<w:r><w:t xml:space="preserve">{txt}</w:t></w:r></w:p>' if txt else f'<w:p><w:r><w:t xml:space="preserve">In the line: </w:t></w:r>{sh}<w:r><w:t xml:space="preserve"> after it.</w:t></w:r></w:p>'
                    x, n = re.subn(rf'<w:p>(?:(?!</w:p>).)*?MARK{i}</w:t></w:r></w:p>', para, x, count=1)
                    assert n == 1, i
                data = x.encode()
            elif it.filename == 'word/header1.xml':
                x = data.decode()
                x, n = re.subn(r'<w:r><w:t>HDRMARK</w:t></w:r>', '<w:r><w:t>Letterhead</w:t></w:r>' + WORD_HDR, x, count=1)
                assert n == 1
                data = x.encode()
            zo.writestr(it, data)
    return base64.b64encode(out.getvalue()).decode()

async def drag(page, a, b, steps=6, shift=False):
    await page.mouse.move(*a)
    if shift: await page.keyboard.down('Shift')
    await page.mouse.down()
    await page.mouse.move(*b, steps=steps)
    await page.mouse.up()
    if shift: await page.keyboard.up('Shift')
    await page.wait_for_timeout(150)

async def main():
    async with async_playwright() as p:
        b, page, errors = await new_page(p, width=1366, height=900, mock_caps=True, gfonts=True)
        await load(page)
        st = lambda js, *a: page.evaluate(js, *a)
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page.wait_for_timeout(500)
        P0 = await st("() => { const r = document.querySelectorAll('#editor p')[1].getBoundingClientRect(); return [r.left, r.top]; }")

        print('1. Insert › Shapes, drawing')
        await page.click('#tab-insert')
        await page.locator('.rb:visible:has(.cap:text-is("Shapes"))').click()
        await page.wait_for_timeout(200)
        n = await st("() => document.querySelectorAll('.pop .shp-btn').length")
        ok(n == 28, 'the Shapes gallery: lines, rectangles, basic shapes, block arrows, callouts and a text box', n)
        await page.locator('.pop .shp-btn[title="Oval"]').click()
        await page.wait_for_timeout(150)
        ok(await st("() => !!document.querySelector('#paper .shp-draw') && document.getElementById('paper').classList.contains('drawing-shape')"), 'a shape picked: the page takes a drag')
        await drag(page, (P0[0] + 330, P0[1] + 20), (P0[0] + 450, P0[1] + 70), shift=True)
        s = await st(SHP)
        ok(len(s) == 1 and s[0]['a'] == {'shp': 'ellipse', 'fill': '#4472c4', 'line': '#2f528f', 'lw': '1', 'va': 'ctr', 'tc': '#ffffff'} and 'img-front' in s[0]['cls'] and s[0]['drawn'], "drawn: Word's blue oval, in front of the text, drawn in its own root", s)
        sz = re.findall(r'width: ([\d.]+)px; height: ([\d.]+)px', s[0]['style'])
        ok(sz and abs(float(sz[0][0]) - float(sz[0][1])) < 0.6 and float(sz[0][0]) >= 119, 'Shift: a circle', s[0]['style'])
        r = await st("() => [__Q.S.selImg && __Q.S.selImg.classList.contains('shp'), __Q.S.tab, document.querySelectorAll('.img-frame .hd').length, !document.querySelector('#paper .shp-draw')]")
        ok(r == [True, 'shape', 9, True], 'selected after drawing: eight handles and the round one, the Shape Format tab', r)
        # Esc cancels; a click (no drag) draws Word's 1-inch shape
        await st("() => __Q.Shapes.draw('rect')")
        await page.keyboard.press('Escape')
        ok(len(await st(SHP)) == 1 and not await st("() => !!document.querySelector('#paper .shp-draw')"), 'Esc: no shape drawn')
        await st("() => __Q.Shapes.draw('rect')")
        await page.mouse.click(P0[0] + 30, P0[1] + 160)
        await page.wait_for_timeout(200)
        s = await st(SHP)
        ok(len(s) == 2 and 'width: 96px; height: 96px' in s[1]['style'] + s[0]['style'], 'a click: a 1-inch square', [x['style'] for x in s])
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        ok(len(await st(SHP)) == 1, 'Undo takes the new shape away')

        print('2. Typing in a shape')
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-shp=ellipse]'))")
        await page.keyboard.type('Seal')
        await page.wait_for_timeout(250)
        r = await st("() => __Q.Story.cur && [__Q.Story.cur.type, __Q.Story.cur.el.textContent, getComputedStyle(__Q.Story.cur.el).color, getComputedStyle(__Q.Story.cur.el.firstElementChild).textAlign]")
        ok(r == ['shape', 'Seal', 'rgb(255, 255, 255)', 'center'], 'a letter typed on a selected shape goes into it (white, centred, as Word)', r)
        await page.keyboard.press('Escape')
        await page.wait_for_timeout(200)
        s = await st(SHP)
        r = await st("() => [__Q.S.selImg === document.querySelector('#editor span.shp[data-shp=ellipse]'), document.querySelector('#editor span.shp[data-shp=ellipse]').shadowRoot.querySelector('.in').textContent, !__Q.Story.cur]")
        ok(s[0]['text'] == '<p style="text-align:center">Seal</p>' and r == [True, 'Seal', True], 'Esc: the text is kept and drawn, the shape stays selected', [s[0]['text'], r])
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        ok((await st(SHP))[0]['text'] in ('', '<p style="text-align:center"><br></p>'), 'one Undo takes the text away', (await st(SHP))[0]['text'])
        await st("() => __Q.redo()")
        await page.wait_for_timeout(200)
        ok('Seal' in (await st(SHP))[0]['text'], 'Redo brings it back')
        # double-click: the text at the click; Enter on a selected shape opens it too
        bx = await st(BOX, 'span.shp[data-shp=ellipse]')
        await page.mouse.dblclick(bx['cx'] + 4, bx['cy'])
        await page.wait_for_timeout(250)
        ok(await st("() => !!__Q.Story.cur && __Q.Story.cur.type === 'shape'"), 'a double-click opens its text')
        await page.keyboard.press('End')
        await page.keyboard.type(' 2026')
        await page.mouse.click(P0[0] + 20, P0[1] + 260)
        await page.wait_for_timeout(250)
        r = await st("() => [!__Q.Story.cur, !__Q.S.selImg, document.querySelector('#editor span.shp[data-shp=ellipse] template').innerHTML]")
        ok(r[0] and r[1] and 'Seal 2026' in r[2], 'a click in the text closes it (the shape not selected)', r)
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-shp=ellipse]'))")
        await page.keyboard.press('Enter')
        await page.wait_for_timeout(200)
        ok(await st("() => !!__Q.Story.cur"), 'Enter on a selected shape opens its text')
        await page.keyboard.press('Escape')
        await page.wait_for_timeout(150)

        print('3. Text boxes')
        await page.click('#tab-insert')
        await page.locator('.rb:visible:has(.cap:text-is("Text Box"))').click()
        await page.wait_for_timeout(150)
        items = await st("() => [...document.querySelectorAll('.pop .mi .lb')].map(e => e.textContent)")
        ok(items == ['Simple Text Box', 'Draw Text Box'], 'Insert › Text Box: Simple, Draw', items)
        await page.locator('.pop .mi', has_text='Draw Text Box').click()
        await page.wait_for_timeout(150)
        await drag(page, (P0[0] + 20, P0[1] + 120), (P0[0] + 260, P0[1] + 150))
        await page.wait_for_timeout(200)
        ok(await st("() => !!__Q.Story.cur && __Q.Story.cur.kind === 'text box'"), 'drawn, a text box opens for typing')
        await page.keyboard.type('Memo No. 12/2026')
        await page.keyboard.press('Enter')
        await page.keyboard.type('Date: 20 October 2026')
        await page.wait_for_timeout(250)
        h0 = await st("() => parseFloat(document.querySelectorAll('#editor span.shp')[1].style.height)")
        for i in range(4):
            await page.keyboard.press('Enter'); await page.keyboard.type(f'Line {i + 3}')
        await page.wait_for_timeout(300)
        h1 = await st("() => parseFloat(document.querySelectorAll('#editor span.shp')[1].style.height)")
        ok(h1 > h0 + 60, 'it grows with its text (fit)', [h0, h1])
        await page.keyboard.press('Escape')
        await page.wait_for_timeout(200)
        s = await st(SHP)
        tb = next(x for x in s if x['a'].get('tb') == '1')
        ok(tb['a'].get('tb') == '1' and tb['a'].get('fill') == '#ffffff' and tb['a'].get('line') == '#000000' and tb['a'].get('lw') == '0.75' and tb['a'].get('fit') == '1' and 'Date: 20 October 2026' in tb['text'], 'a text box: white, a thin black line, fits its text', tb)
        # a click inside a text box types there; one on its edge selects it
        await st("() => __Q.Images.deselect()")
        bx = await st(BOX, 'span.shp[data-tb]')
        await page.mouse.click(bx['cx'], bx['cy'])
        await page.wait_for_timeout(200)
        ok(await st("() => !!__Q.Story.cur && __Q.Story.cur.shape === document.querySelectorAll('#editor span.shp')[1]"), 'a click inside a text box: typing in it')
        await page.keyboard.press('Escape')
        await st("() => __Q.Images.deselect()")
        await page.mouse.click(bx['cx'] - bx['w'] / 2 + 2, bx['cy'])
        await page.wait_for_timeout(200)
        ok(await st("() => !__Q.Story.cur && __Q.S.selImg === document.querySelectorAll('#editor span.shp')[1]"), 'a click on its edge: selected')
        await page.keyboard.press('Delete')
        await page.wait_for_timeout(150)
        ok(len(await st(SHP)) == 1, 'Delete removes it')
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        ok(len(await st(SHP)) == 2 and all(x['drawn'] for x in await st(SHP)), 'Undo brings it back, drawn')

        print('4. Handles, keys')
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-shp=ellipse]'))")
        hd = await st("() => { const r = document.querySelector('.img-frame .hd.se').getBoundingClientRect(); return [r.left + 5, r.top + 5]; }")
        await drag(page, hd, (hd[0] + 50, hd[1] + 12), shift=True)
        s = [x for x in await st(SHP) if x['a'].get('shp') == 'ellipse']
        w, hh = [float(x) for x in re.findall(r'width: ([\d.]+)px; height: ([\d.]+)px', s[0]['style'])[0]]
        ok(abs(w - hh) < 1 and w > 150, 'a corner with Shift keeps the proportions', [w, hh])
        rh = await st("() => { const r = document.querySelector('.img-frame .hd.rot').getBoundingClientRect(); return [r.left + 7, r.top + 7]; }")
        c = await st(BOX, 'span.shp[data-shp=ellipse]')
        await drag(page, rh, (c['cx'] + 100, c['cy'] - 50), shift=True)
        rot = await st("() => +document.querySelector('#editor span.shp[data-shp=ellipse]').dataset.rot")
        ok(rot % 15 == 0 and 30 <= rot <= 90, 'the round handle turns it; Shift: 15° steps', rot)
        fr = await st("() => getComputedStyle(document.querySelector('.img-frame')).transform !== 'none'")
        ok(fr, 'the frame turns with it')
        # the west side of the turned shape: the east side stays where it was
        b0 = await st(BOX, 'span.shp[data-shp=ellipse]')
        east = lambda bb: (bb['cx'] + math.cos(math.radians(bb['rot'])) * bb['w'] / 2, bb['cy'] + math.sin(math.radians(bb['rot'])) * bb['w'] / 2)
        e0 = east(b0)
        wh = await st("() => { const r = document.querySelector('.img-frame .hd.w').getBoundingClientRect(); return [(r.left + r.right) / 2, (r.top + r.bottom) / 2]; }")
        await drag(page, wh, (wh[0] - 30, wh[1] - 10))
        b1 = await st(BOX, 'span.shp[data-shp=ellipse]')
        e1 = east(b1)
        ok(b1['w'] > b0['w'] + 10 and abs(e1[0] - e0[0]) < 1.5 and abs(e1[1] - e0[1]) < 1.5, 'a turned shape widened from its left side keeps its right side in place', [e0, e1, b0['w'], b1['w']])
        l0 = await st("() => parseFloat(document.querySelector('#editor span.shp[data-shp=ellipse]').style.left)")
        await page.keyboard.press('ArrowRight')
        await page.keyboard.press('Shift+ArrowDown')
        r = await st("() => { const s = document.querySelector('#editor span.shp[data-shp=ellipse]'); return [parseFloat(s.style.left), parseFloat(s.style.top)]; }")
        ok(abs(r[0] - l0 - 2) < 0.2, 'arrow keys move it (2 px; Shift 10)', [l0, r])
        # a line: a handle at each end; Shift: 45°
        await st("() => __Q.Shapes.draw('line:end')")
        await drag(page, (P0[0] + 40, P0[1] + 330), (P0[0] + 200, P0[1] + 350))
        r = await st("() => [document.querySelectorAll('.img-frame .hd').length, !!document.querySelector('.img-frame .hd.p1'), !!document.querySelector('.img-frame .hd.rot'), document.querySelector('.img-frame').classList.contains('line')]")
        ok(r == [2, True, False, True], 'a line: a handle at each end, no box, no round handle', r)
        p2 = await st("() => { const r = document.querySelector('.img-frame .hd.p2').getBoundingClientRect(); return [(r.left + r.right) / 2, (r.top + r.bottom) / 2]; }")
        await drag(page, p2, (p2[0] - 40, p2[1] + 70), shift=True)
        LINE = lambda lst: next(x for x in lst if x['a'].get('shp') == 'line')
        ln = LINE(await st(SHP))
        w, hh = [float(x) for x in re.findall(r'width: ([\d.]+)px; height: ([\d.]+)px', ln['style'])[0]]
        ok(abs(w - hh) < 1 and ln['a'].get('head') == 'end', 'a line\'s end dragged with Shift: at 45°', ln)
        p2 = await st("() => { const r = document.querySelector('.img-frame .hd.p2').getBoundingClientRect(); return [(r.left + r.right) / 2, (r.top + r.bottom) / 2]; }")
        await drag(page, p2, (p2[0] - 300, p2[1] - 20))
        ln = LINE(await st(SHP))
        ok(ln['a'].get('flip') == 'h', 'dragged past its other end: the line points the other way (flipped)', ln['a'])

        print('5. Shape Format')
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-shp=ellipse]'))")
        await page.click('#tab-shape')
        await page.locator('.rb:visible:has(.cap:text-is("Styles"))').click()
        await page.wait_for_timeout(150)
        ok(await st("() => document.querySelectorAll('.pop .shp-pre').length") == 16, 'Shape Styles: 16 looks')
        await page.locator('.pop .shp-pre[title="Light yellow"]').click()
        await page.wait_for_timeout(150)
        a = next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a']
        ok(a.get('fill') == '#fff2cc' and a.get('line') == '#bf9000' and a.get('tc') == '#3f3000', 'a style: fill, line and text colour', a)
        await page.locator('.rb:visible:has(.cap:text-is("Shape Fill"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .sw[title="#00b0f0"]').click()
        await page.wait_for_timeout(150)
        await page.locator('.rb:visible:has(.cap:text-is("Shape Outline"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .shp-lb[title="3 pt"]').click()
        await page.wait_for_timeout(150)
        await page.locator('.rb:visible:has(.cap:text-is("Shape Outline"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .shp-lb[title="Dash"]').click()
        await page.wait_for_timeout(150)
        a = next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a']
        ok(a.get('fill') == '#00b0f0' and a.get('lw') == '3' and a.get('dash') == 'dash', 'Shape Fill, Shape Outline (weight, dashes)', a)
        stroke = await st("() => { const p = document.querySelector('#editor span.shp[data-shp=ellipse]').shadowRoot.querySelector('path'); return [p.getAttribute('stroke-width'), p.getAttribute('stroke-dasharray')]; }")
        ok(stroke == ['4', '16 12'], 'drawn with the new outline', stroke)
        await page.locator('.rb:visible:has(.cap:text-is("Align Text"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .mi', has_text='Bottom').click()
        await page.wait_for_timeout(150)
        ok(next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a'].get('va') == 'b' and await st("() => document.querySelector('#editor span.shp[data-shp=ellipse]').shadowRoot.querySelector('.tx').classList.contains('b')"), 'Align Text: bottom')
        await page.locator('#shpH').fill('1.5"')
        await page.locator('#shpH').press('Enter')
        await page.wait_for_timeout(200)
        ok('height: 144px' in next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['style'], 'Height box: 1.5" = 144 px', next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['style'])
        rot0 = +next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a'].get('rot', 0) if False else float(next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a'].get('rot', '0'))
        await page.locator('.rb:visible:has(.cap:text-is("Rotate"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .mi', has_text='Rotate Right 90°').click()
        await page.wait_for_timeout(150)
        await page.locator('.rb:visible:has(.cap:text-is("Rotate"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .mi', has_text='Flip Horizontal').click()
        await page.wait_for_timeout(150)
        a = next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a']
        ok(abs(float(a.get('rot', 0)) - (rot0 + 90) % 360) < 0.01 and a.get('flip') == 'h', 'Rotate Right 90°, Flip Horizontal', [rot0, a])
        # wrapping
        async def wrap(label):
            await page.locator('.rb:visible:has(.cap:text-is("Wrap Text"))').click()
            await page.wait_for_timeout(150)
            await page.locator('.pop .mi', has_text=label).click()
            await page.wait_for_timeout(250)
            return next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')
        s0 = await wrap('In Line with Text')
        ok(s0['cls'] == 'shp' and 'left' not in s0['style'], 'In Line with Text: in the line', s0)
        y = await st("() => { const s = document.querySelector('#editor span.shp[data-shp=ellipse]'); const p = s.closest('p'); return [getComputedStyle(s).position, p.getBoundingClientRect().height > 100]; }")
        ok(y == ['relative', True], 'its line is as tall as it is', y)
        s0 = await wrap('Square, shape on the left')
        ok('img-left' in s0['cls'] and await st("() => getComputedStyle(document.querySelector('#editor span.shp[data-shp=ellipse]')).float") == 'left', 'Square: text beside it', s0['cls'])
        s0 = await wrap('Behind Text')
        ok('img-behind' in s0['cls'] and 'left:' in s0['style'] and await st("() => getComputedStyle(document.querySelector('#editor span.shp[data-shp=ellipse]')).zIndex") == '-1', 'Behind Text', s0)
        s0 = await wrap('In Front of Text')
        await page.locator('.rb:visible:has(.cap:text-is("Center"))').first.click()
        await page.wait_for_timeout(200)
        r = await st("() => { const s = document.querySelector('#editor span.shp[data-shp=ellipse]'); const r = s.getBoundingClientRect(); const ed = document.querySelector('#editor').getBoundingClientRect(); const g = __Q.geom(), z = __Q.S.zoom; return [(r.left + r.right) / 2 - ed.left, (g.ml + g.cw / 2) * z]; }")
        ok(abs(r[0] - r[1]) < 1.5, 'Center: in the middle of the column', r)
        # Format Shape dialog: one step to undo
        await page.locator('.rgroup[aria-label="Shape Styles"] .launcher').click()
        await page.wait_for_timeout(200)
        dlg = page.locator('.dialog')
        ok(await dlg.locator('h2').text_content() == 'Format Shape', 'the launcher opens Format Shape')
        await dlg.locator('input[type=checkbox]').first.check()
        await dlg.locator('input[type=number]').first.fill('1.5')
        await dlg.locator('select').nth(1).select_option('t')
        await page.click('.dialog footer .btn.primary')
        await page.wait_for_timeout(200)
        a = next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a']
        ok(a.get('fill') == 'none' and a.get('lw') == '1.5' and a.get('va') == 't', 'Format Shape: no fill, 1.5 pt, text at the top', a)
        await st("() => __Q.undo()")
        await page.wait_for_timeout(200)
        a = next(x for x in await st(SHP) if x['a'].get('shp') == 'ellipse')['a']
        ok(a.get('fill') == '#00b0f0' and a.get('lw') == '3' and a.get('va') == 'b', 'one Undo puts it all back', a)
        # fit turned on: exactly the text's height
        TBQ = "document.querySelector('#editor span.shp[data-tb]')"
        await st(f"() => {{ const t = {TBQ}; __Q.Images.select(t); __Q.Shapes.toggleFit(); __Q.Shapes.setSize(null, 400); }}")
        await page.wait_for_timeout(150)
        hb = await st(f"() => parseFloat({TBQ}.style.height)")
        await st("() => __Q.Shapes.toggleFit()")
        await page.wait_for_timeout(200)
        ha = await st(f"() => parseFloat({TBQ}.style.height)")
        ok(hb == 400 and 80 < ha < 200, 'Resize Shape to Fit Text turned on: it takes its text\'s height', [hb, ha])
        # arrows on a line
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-shp=line]'))")
        await page.locator('.rb:visible:has(.cap:text-is("Shape Outline"))').click()
        await page.wait_for_timeout(150)
        await page.locator('.pop .shp-lb[title="Arrows at both ends"]').click()
        await page.wait_for_timeout(150)
        r = await st("() => { const l = document.querySelector('#editor span.shp[data-shp=line]'); return [l.dataset.head, l.shadowRoot.querySelectorAll('svg path').length]; }")
        ok(r == ['both', 4], 'Shape Outline › Arrows: both ends (two heads drawn)', r)

        print('6. Copy, paste, cut; saved')
        await st("() => __Q.Images.select(document.querySelector('#editor span.shp[data-tb]'))")
        await page.keyboard.press('Control+c')
        await st(CARET_END, 4)
        await page.keyboard.press('Control+v')
        await page.wait_for_timeout(400)
        s = await st(SHP)
        tbs = [x for x in s if x['a'].get('tb') == '1']
        ok(len(s) == 4 and len(tbs) == 2 and tbs[0]['a'] == tbs[1]['a'] and tbs[0]['text'] == tbs[1]['text'] and all(x['drawn'] for x in tbs), 'Ctrl+C, Ctrl+V: the same text box again, drawn', [len(s), tbs])
        await st("() => __Q.Images.select(document.querySelectorAll('#editor span.shp[data-tb]')[1])")
        await page.keyboard.press('Control+x')
        await page.wait_for_timeout(300)
        ok(len(await st(SHP)) == 3, 'Ctrl+X takes it out')
        r = await st("() => [...document.querySelectorAll('#editor p')].map(p => p.textContent).slice(-2)")
        ok(r == ['Copy to all sections.', 'End of order.'], 'and nothing else (the browser\'s own delete took a letter instead)', r)
        r = await st("""async () => { await __Q.Saver.flush(); __Q.Saver.mark(); await __Q.Saver.flush(); const id = __Q.S.doc.id; const back = await __Q.CloudStore.load(id); const d = document.createElement('div'); d.innerHTML = back.html; return [d.querySelectorAll('span.shp').length, [...d.querySelectorAll('span.shp template')].map(t => t.innerHTML).join(' | ')]; }""")
        ok(r[0] == 3 and 'Seal 2026' in r[1] and 'Memo No. 12/2026' in r[1], 'saved to the account: three shapes with their text', r)
        await st("""async () => { const id = __Q.S.doc.id; const back = await __Q.CloudStore.load(id); __Q.openDoc(back); }""")
        await page.wait_for_timeout(600)
        s = await st(SHP)
        ok(len(s) == 3 and all(x['drawn'] for x in s) and 'Seal 2026' in await st("() => document.querySelector('#editor span.shp[data-shp=ellipse]').shadowRoot.textContent"), 'opened again: drawn, with their text', [x['drawn'] for x in s])
        await page.screenshot(path='shots/shp1.png')

        print('7. PDF')
        os.makedirs('shots', exist_ok=True)
        open('shots/shp1.pdf', 'wb').write(base64.b64decode(await st(GET_PDF)))
        txt = pdf_text('shots/shp1.pdf')
        # the oval is turned: its words come out of the PDF in pieces, in the order they're drawn
        raw = re.sub(r'\s+', '', subprocess.run(['pdftotext', '-raw', 'shots/shp1.pdf', '-'], capture_output=True, text=True).stdout)
        ok('Seal2026' in raw and 'Memo No. 12/2026' in txt and 'Line 6' in txt, 'the vector PDF has the shapes\' text', [txt[:600], raw[:300]])
        img = page_png('shots/shp1.pdf', 72)
        # the oval's fill where nothing is written: its centre, a little up (its text is at the bottom)
        c = await st("""() => { const s = document.querySelector('#editor span.shp[data-shp=ellipse]'), sh = document.getElementById('sheets').firstElementChild.getBoundingClientRect(); const r = s.getBoundingClientRect(), z = __Q.S.zoom; return [((r.left + r.right) / 2 - sh.left) / z * 0.75, ((r.top + r.bottom) / 2 - sh.top) / z * 0.75 - 12]; }""")
        px = img.getpixel((int(c[0]), int(c[1])))
        ok(abs(px[0] - 0x00) < 40 and abs(px[1] - 0xb0) < 40 and abs(px[2] - 0xf0) < 40, 'the PDF draws its fill (sky blue)', [c, px])
        canvas = await st("""async () => { const P = await __Q.collectPaint(); const cv = __Q.paintPage(P, 0, 1); const s = document.querySelector('#editor span.shp[data-shp=ellipse]'), sh = document.getElementById('sheets').firstElementChild.getBoundingClientRect(); const r = s.getBoundingClientRect(), z = __Q.S.zoom; const x = Math.round(((r.left + r.right) / 2 - sh.left) / z), y = Math.round(((r.top + r.bottom) / 2 - sh.top) / z - 16); return Array.from(cv.getContext('2d').getImageData(x, y, 1, 1).data).slice(0, 3); }""")
        ok(abs(canvas[0] - 0x00) < 40 and abs(canvas[1] - 0xb0) < 40 and abs(canvas[2] - 0xf0) < 40, 'the page painter too (print preview, picture PDF)', canvas)

        print('8. The Word file')
        b64 = await st(GET_DOCX)
        z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
        doc = z.read('word/document.xml').decode()
        ok(doc.count('<wps:wsp>') == 3 and doc.count('<mc:Choice Requires="wps">') == 3 and 'xmlns:wps=' in doc[:2000], 'three Word 2010 shapes (wps) in mc:AlternateContent', doc.count('<wps:wsp>'))
        ok('<a:prstGeom prst="ellipse">' in doc and '<a:ln w="38100"><a:solidFill><a:srgbClr val="BF9000"/></a:solidFill><a:prstDash val="dash"/>' in doc and '<a:solidFill><a:srgbClr val="00B0F0"/></a:solidFill>' in doc, 'the oval: its fill, a 3 pt dashed line', re.findall(r'<wps:spPr>.*?</wps:spPr>', doc)[:1])
        ok(re.search(r'<a:xfrm rot="\d+" flipH="1">', doc) is not None and 'anchor="b"' in doc, 'turned and flipped; text at the bottom', re.findall(r'<a:xfrm[^>]*>', doc))
        ok('<wps:cNvSpPr txBox="1"/>' in doc and '<a:spAutoFit/>' in doc and 'Memo No. 12/2026</w:t>' in doc, 'the text box: txBox, fitting its text, its text in txbxContent')
        ok('<a:headEnd type="triangle"/><a:tailEnd type="triangle"/>' in doc and '<wps:cNvCnPr/>' in doc and '<a:prstGeom prst="line">' in doc, 'the line with arrows at both ends')
        ok('<a:fontRef idx="minor"><a:srgbClr val="3F3000"/></a:fontRef>' in doc, 'a shape\'s own text colour as the style\'s font reference')
        ok(doc.count('<wp:wrapNone/>') == 3 and doc.count('relativeFrom="paragraph"') >= 3, 'in front of the text: wrapNone, from the column and the paragraph')
        await st(OPEN_DOCX, b64)
        await page.wait_for_timeout(700)
        s = await st(SHP)
        s.sort(key=lambda x: {'ellipse': 0, 'rect': 1, 'line': 2}.get(x['a'].get('shp'), 3))
        ok(len(s) == 3 and s[0]['a'].get('shp') == 'ellipse' and s[0]['a'].get('fill') == '#00b0f0' and s[0]['a'].get('line') == '#bf9000' and s[0]['a'].get('lw') == '3' and s[0]['a'].get('dash') == 'dash' and s[0]['a'].get('flip') == 'h' and s[0]['a'].get('va') == 'b' and s[0]['a'].get('tc') == '#3f3000' and 'img-front' in s[0]['cls'] and s[0]['drawn'], 'opened again: the oval as it was', s[0])
        ok(s[1]['a'].get('tb') == '1' and s[1]['a'].get('fit') == '1' and 'Date: 20 October 2026' in s[1]['text'] and s[2]['a'].get('head') == 'both' and s[2]['a'].get('shp') == 'line', 'the text box and the line too', [s[1]['a'], s[2]['a']])
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, 'shapes.docx'); open(src, 'wb').write(base64.b64decode(b64))
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', td, src], capture_output=True, timeout=120)
            lo = os.path.join(td, 'shapes.pdf')
            t = pdf_text(lo) if os.path.exists(lo) else ''
            ok('Seal 2026' in t and 'Memo No. 12/2026' in t, 'LibreOffice opens it with the shapes\' text', t[:500])
            if os.path.exists(lo):
                im = page_png(lo, 36)
                cols = im.getcolors(1 << 20)
                sky = sum(n for n, (r_, g_, b_) in cols if abs(r_ - 0) < 40 and abs(g_ - 0xb0) < 40 and abs(b_ - 0xf0) < 40)
                ok(sky > 150, 'and draws the oval\'s fill', sky)

        print('9. Shapes made by Word')
        await st(OPEN_DOCX, word_fixture())
        await page.wait_for_timeout(900)
        s = await st(SHP)
        ok(len(s) == 6, 'all of them: a text box, an oval, a group of two, an old VML box, one in the line', [x['a'].get('shp') for x in s])
        bx = next((x for x in s if 'Memo No. 45/2026' in x['text']), None)
        ok(bx and bx['a'].get('tb') == '1' and 'img-right' in bx['cls'] and bx['a'].get('fill') == '#ffffff' and bx['a'].get('line') == '#000000' and bx['a'].get('lw') == '0.75' and bx['a'].get('fit') == '1', 'Word 2016\'s text box (square wrap at the right): to the right, the text beside it', bx)
        ov = next((x for x in s if 'Seal' in x['text']), None)
        ok(ov and ov['a'].get('shp') == 'ellipse' and ov['a'].get('fill') == '#4f81bd' and ov['a'].get('line') == '#385d8a' and ov['a'].get('lw') == '1' and ov['a'].get('tc') == '#ffffff' and ov['a'].get('rot') == '30' and ov['a'].get('va') == 'ctr' and 'text-align:center' in ov['text'], 'Word 365\'s oval: colours from its style and the file\'s theme, turned 30°, white centred text', ov)
        pos = await st("() => { const s = [...document.querySelectorAll('#editor span.shp')].find(x => x.dataset.shp === 'ellipse'); return [parseFloat(s.style.left), parseFloat(s.style.top)]; }")
        ok(abs(pos[0] - 192) < 1.5 and abs(pos[1] - 96) < 1.5, 'placed as in Word: 2" from the column, 1" below its paragraph', pos)
        grp = [x for x in s if x['a'].get('fill') == '#ffc000' or (x['a'].get('shp') == 'line' and x['a'].get('line') == '#c00000')]
        ok(len(grp) == 2 and any(x['a'].get('head') == 'end' and x['a'].get('lw') == '1.5' for x in grp) and all('img-front' in x['cls'] for x in grp), 'the group: its box and its red arrow, each on its own', grp)
        old = next((x for x in s if 'Old Word box' in x['text']), None)
        ok(old and old['a'].get('tb') == '1' and old['a'].get('line') == '#c00000' and old['a'].get('lw') == '1.5' and old['a'].get('fit') == '1' and old['a'].get('ins') == '5.4 3.6 5.4 3.6' and 'img-front' in old['cls'], 'the older (VML) text box: its line, insets, fitting', old)
        r = await st("() => { const s = [...document.querySelectorAll('#editor span.shp')].find(x => /Old Word box/.test((x.querySelector('template') || { innerHTML: '' }).innerHTML)); const g = __Q.geom(); return [parseFloat(s.style.left) + parseFloat(s.style.width), g.cw]; }")
        ok(abs(r[0] - r[1]) < 1.5, 'at the right of the margin, as Word put it', r)
        il = next((x for x in s if x['a'].get('fill') == '#00b050'), None)
        ok(il and il['cls'] == 'shp' and 'width: 48px; height: 24px' in il['style'] and il['a'].get('line') == 'none', 'the one in the line: in the line, its size, no outline', il)
        hdr = await st("() => { const s = document.querySelector('#sheets .sheet .hfx span.shp'); return s && [!!s.shadowRoot, s.shadowRoot && s.shadowRoot.textContent.trim(), s.dataset.fill, s.dataset.line]; }")
        ok(hdr and hdr[0] and hdr[1] == 'Ref: 12/2026' and hdr[2] == 'none' and hdr[3] == '#1f3a5f', 'the header\'s text box, drawn on the page', hdr)
        open('shots/shp1-word.pdf', 'wb').write(base64.b64decode(await st(GET_PDF)))
        t = pdf_text('shots/shp1-word.pdf')
        ok('Ref: 12/2026' in t and 'Memo No. 45/2026' in t and 'Old Word box' in t, 'the PDF has them all, the header\'s too', t[:500])
        await page.screenshot(path='shots/shp1-word.png')

        print('10. Web page, Markdown, text')
        html = await st("() => __Q.toStandaloneHtml()")
        ok(html.count('shadowrootmode="open"') == 7, 'the web page: each shape drawn by the page itself (declarative shadow root)', html.count('shadowrootmode'))
        p2 = await b.new_page()
        await p2.set_content(html)
        await p2.wait_for_timeout(300)
        r = await p2.evaluate("() => [...document.querySelectorAll('span.shp')].map(s => !!s.shadowRoot && !!s.shadowRoot.querySelector('svg path'))")
        ok(len(r) == 7 and all(r), 'opened in a browser, every shape is drawn', r)
        await p2.close()
        md = await st("() => __Q.toMarkdown()")
        tx = await st("() => __Q.toPlainText()")
        ok('Memo No. 45/2026' in md and 'Seal' in md and 'Old Word box' in tx, 'Markdown and plain text carry the shapes\' text', md[:400])

        print('11. Read-only, phone')
        await st("() => { __Q.S.readOnly = true; }")
        await st("() => __Q.Shapes.draw('rect')")
        ok(not await st("() => !!document.querySelector('#paper .shp-draw')"), 'read-only: nothing to draw')
        await st("() => { __Q.Images.select(document.querySelector('#editor span.shp')); }")
        await page.keyboard.type('x')
        await page.wait_for_timeout(150)
        ok(not await st("() => !!__Q.Story.cur"), 'read-only: a letter doesn\'t open a shape\'s text')
        await st("() => { __Q.S.readOnly = false; __Q.Images.deselect(); }")

        print('12. Backspace and Delete beside a shape')
        KBOX = '<span class="shp img-front" contenteditable="false" data-shp="rect" data-tb="1" data-fill="#ffffff" data-line="#000000" data-lw="0.75" data-va="t" style="width:150px;height:40px;left:260px;top:4px;"><template><p>Ref: 7</p></template></span>'
        KEYS = f'<p>The office will remain closed.</p><p>{KBOX}All staff.</p><p>Duty officers.</p>'
        KST = """() => { const ps = [...document.querySelectorAll('#editor > p')]; const s = document.querySelector('#editor span.shp'); const r = s && s.getBoundingClientRect(); return [ps.map(p => p.textContent), document.querySelectorAll('#editor span.shp').length, r ? [Math.round(r.left), Math.round(r.top)] : null, s ? ps.indexOf(s.parentElement) : -1]; }"""
        CARET = """([i, end]) => { const p = document.querySelectorAll('#editor > p')[i]; const t = [...p.childNodes].filter(n => n.nodeType === 3); const n = end ? t[t.length - 1] : t[0]; const r = document.createRange(); r.setStart(n, end ? n.length : 0); __Q.setRange(r); document.getElementById('editor').focus(); }"""
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Keys', html: h }))", KEYS)
        await page.wait_for_timeout(300)
        k0 = await st(KST)
        await st(CARET, [1, False])
        await page.keyboard.press('Backspace')
        await page.wait_for_timeout(200)
        k1 = await st(KST)
        ok(k1[0] == ['The office will remain closed.All staff.', 'Duty officers.'] and k1[1] == 1 and k1[2] == k0[2] and k1[3] == 0, 'Backspace at the start of the text after a text box: the paragraphs join, the box stays where it was (as Word)', [k0, k1])
        await page.keyboard.type('X')
        await page.wait_for_timeout(150)
        ok((await st(KST))[0][0] == 'The office will remain closed.XAll staff.', 'and the caret is at the join')
        await st("() => __Q.undo()"); await page.wait_for_timeout(150)
        await st("() => __Q.undo()"); await page.wait_for_timeout(200)
        ok(await st(KST) == k0, 'Undo puts the paragraph and its box back', await st(KST))
        await st(CARET, [0, True])
        await page.keyboard.press('Delete')
        await page.wait_for_timeout(200)
        k2 = await st(KST)
        ok(k2[0] == k1[0] and k2[1] == 1 and k2[2] == k0[2], 'Delete at the end of the paragraph before: the same', k2)
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Keys', html: h }))", KEYS.replace('<p>The office will remain closed.</p>', ''))
        await page.wait_for_timeout(300)
        await st(CARET, [0, False])
        await page.keyboard.press('Backspace')
        await page.wait_for_timeout(200)
        ok(await st("() => document.querySelectorAll('#editor span.shp').length") == 1, 'at the start of the document Backspace leaves the box alone')
        await st("h => __Q.openDoc(__Q.newDocModel({ title: 'Keys', html: h }))", f'<p>The office will remain closed.</p><p>All {KBOX}staff.</p><p>Duty officers.</p>')
        await page.wait_for_timeout(300)
        await st("""() => { const p = document.querySelectorAll('#editor > p')[1]; const t = [...p.childNodes].filter(n => n.nodeType === 3)[1]; const r = document.createRange(); r.setStart(t, 0); __Q.setRange(r); }""")
        await page.keyboard.press('Backspace')
        await page.wait_for_timeout(200)
        k3 = await st(KST)
        ok(k3[0][1] == 'Allstaff.' and k3[1] == 1, 'a box in the middle of a line: Backspace after it deletes the space before it', k3)
        errs = [e for e in errors if 'favicon' not in e[1]]
        ok(not errs, 'no errors in the console', errs[:5])
        await b.close()

        b2, page2, errors2 = await new_page(p, width=400, height=860, mock_caps=True, gfonts=True)
        await load(page2)
        await page2.evaluate("h => __Q.openDoc(__Q.newDocModel({ title: 'Office order', html: h }))", HTML)
        await page2.wait_for_timeout(400)
        await page2.click('#tab-insert')
        await page2.locator('.rb:visible:has(.cap:text-is("Shapes"))').click()
        await page2.wait_for_timeout(200)
        r = await page2.evaluate("() => { const m = document.querySelector('.pop'); const r = m.getBoundingClientRect(); return [r.left >= 0, r.right <= innerWidth, document.documentElement.scrollWidth <= innerWidth]; }")
        ok(r == [True, True, True], 'phone: the Shapes gallery fits the screen', r)
        await page2.locator('.pop .shp-btn[title="Rounded Rectangle"]').click()
        await page2.wait_for_timeout(150)
        pr = await page2.evaluate("() => { const r = document.querySelectorAll('#editor p')[1].getBoundingClientRect(); return [r.left, r.top]; }")
        await page2.mouse.click(pr[0] + 40, pr[1] + 30)
        await page2.wait_for_timeout(300)
        r = await page2.evaluate("() => [document.querySelectorAll('#editor span.shp').length, __Q.S.tab, !document.getElementById('tab-shape').hidden, document.documentElement.scrollWidth <= innerWidth]")
        ok(r == [1, 'shape', True, True], 'phone: a tap draws one; Shape Format shows', r)
        await page2.screenshot(path='shots/shp1-phone.png')
        for e in errors2: print('ERR(phone)', e)
        await b2.close()
    print('\nFAILED: ' + ', '.join(fails) if fails else '\nall passed')

asyncio.run(main())
